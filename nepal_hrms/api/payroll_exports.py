# Copyright (c) 2026, Yarsa Labs Pvt. Ltd. and contributors
# For license information, please see license.txt

"""Excel downloads from a Payroll Entry: bank transfer, SSF and CIT.

- Net pay as a connectIPS Corporate Pay "CIPS BULK TRANSFER" sheet, optionally
  less some earnings (Festival Allowance, Bonus ...) that are paid separately.
- Those earnings on their own, in the same bank sheet.
- SSF and CIT contribution lists.

Components are offered by root name: "Festival Allowance" covers both its
"- Salary" and "- Wages" variants (see custom_code/payroll/component_variant.py).

Amounts are the slip's own figures, rounded to 2 decimals. Net pay is net_pay,
or rounded_total when Payroll Settings rounds it -- the figure HRMS's bank entry
uses. An earning left out of net pay is taken off at the same 2-decimal amount
its own file pays, so the two files add up to the slip exactly.
"""

import os
import re
from collections import defaultdict
from io import BytesIO

import frappe
from frappe import _
from frappe.utils import cint, flt, getdate

from nepal_hrms.custom_code.payroll.component_variant import component_variant

TEMPLATE = os.path.join(
	os.path.dirname(os.path.dirname(__file__)),
	"custom_code",
	"payroll",
	"templates",
	"cips_bulk_transfer.xlsx",
)
BANK_FIRST_ROW = 5  # rows 1-4 of the template are its notes and headers
BANK_MAX_ROWS = 1000  # "Allowed transaction count is 1000"
REMARKS_MAX = 30
ACCOUNT_NAME_MAX = 140

SSF_EMPLOYEE_ROOT = "Employee's Contribution SSF"
SSF_EMPLOYER_ROOT = "Employer's Contribution SSF"
CIT_ROOT = "CIT"

KINDS = ("net_pay", "component_net_pay", "components", "ssf", "cit", "jv")


# --------------------------------------------------------------------------
# Pure helpers (unit-tested without a site)


def component_root(salary_component):
	"""'Festival Allowance - Wages' -> 'Festival Allowance'; shared names as is."""
	name = (salary_component or "").strip()
	variant = component_variant(name)
	return name[: -len(f" - {variant}")] if variant else name


def normalize_bank_name(name):
	"""'Nabil Bank Limited' and 'Nabil Bank' must find the same Bank record."""
	words = re.sub(r"[^a-z0-9 ]", " ", (name or "").lower()).split()
	return " ".join(w for w in words if w not in ("limited", "ltd"))


def bank_code_for(bank_name, codes_by_name):
	return codes_by_name.get(normalize_bank_name(bank_name)) or ""


def counted_in_totals(d):
	"""Whether the slip's gross / total deduction includes this row.

	A "Calculation Only" row (do_not_include_in_total) -- the Leave and Late
	Deduction in earned-salary mode -- has already been taken out of earnings;
	counting it again would pay less than the slip says.
	"""
	return not cint(d.get("do_not_include_in_total")) and not cint(d.get("statistical_component"))


def component_amounts(details, parentfield="earnings"):
	"""{(slip, root): amount} for one table, each rounded once to 2 decimals."""
	totals = defaultdict(float)
	for d in details:
		if d.parentfield == parentfield and counted_in_totals(d):
			totals[(d.parent, component_root(d.salary_component))] += flt(d.amount)
	return {key: flt(value, 2) for key, value in totals.items()}


def offset_roots(details):
	"""Earnings cancelled by a deduction of the same name + " Deduction".

	"Employer's Contribution SSF" is added to gross only to be deducted back as
	"Employer's Contribution SSF Deduction" and paid to the fund, never to the
	employee. Leaving it out of net pay, or paying it to the employee's account,
	is always a mistake, so it is not offered.
	"""
	deductions = {
		component_root(d.salary_component) for d in details if d.parentfield == "deductions"
	}
	return {
		component_root(d.salary_component)
		for d in details
		if d.parentfield == "earnings"
		and f"{component_root(d.salary_component)} Deduction" in deductions
	}


def offset_deduction_roots(details):
	"""The deduction halves of offset_roots: hidden with them, as a pair."""
	return {f"{root} Deduction" for root in offset_roots(details)}


def component_net_amounts(slips, details, earnings, deductions):
	"""{slip: ticked earnings - ticked deductions}, from counted rows only.

	With every earning and deduction ticked this is the slip's net pay: the
	offset pair is left out on both sides and cancels anyway.
	"""
	paid = component_amounts(details, "earnings")
	taken = component_amounts(details, "deductions")
	return {
		s.name: flt(
			sum(paid.get((s.name, root), 0.0) for root in earnings)
			- sum(taken.get((s.name, root), 0.0) for root in deductions),
			2,
		)
		for s in slips
	}


def amounts_by_root(details, root):
	"""{slip: amount} for one component root, earning or deduction alike.

	"Employer's Contribution SSF" is the earning pair; its offsetting deduction
	is named "... Deduction" and so is not counted twice.
	"""
	totals = defaultdict(float)
	for d in details:
		if component_root(d.salary_component) == root:
			totals[d.parent] += flt(d.amount)
	return {slip: flt(value, 2) for slip, value in totals.items()}


def bank_rows(slips, amounts, codes_by_name, remarks):
	"""Rows for the bank sheet, and the slips that could not go on it.

	amounts maps slip name -> amount to pay. A slip with nothing to pay, or with
	no bank account (paid by cash or cheque), is left out; the download dialog
	names the latter beforehand. A bank-paid slip that cannot be written in full
	is reported, never written with a blank required column.
	"""
	rows, problems = [], []
	for slip in slips:
		amount = flt(amounts.get(slip.name), 2)
		account = (slip.bank_account_no or "").strip()
		if amount < 0:
			problems.append((slip, _("Amount to pay is negative ({0})").format(amount)))
			continue
		if amount == 0 or not account:
			continue
		code = bank_code_for(slip.bank_name, codes_by_name)
		if not code:
			problems.append(
				(slip, _("No NCHL bank code for bank '{0}'").format(slip.bank_name or ""))
			)
			continue
		rows.append(
			[
				amount,
				code,
				account,
				(slip.employee_name or "")[:ACCOUNT_NAME_MAX],
				remarks[:REMARKS_MAX],
				slip.employee,
			]
		)
	return rows, problems


# --------------------------------------------------------------------------
# Data


def _load(payroll_entry):
	pe = frappe.get_doc("Payroll Entry", payroll_entry)
	pe.check_permission("read")
	if not frappe.has_permission("Salary Slip", "read"):
		frappe.throw(_("Not permitted to read Salary Slips"), frappe.PermissionError)

	slips = frappe.get_all(
		"Salary Slip",
		filters={"payroll_entry": pe.name, "docstatus": ["<", 2]},
		fields=[
			"name", "employee", "employee_name", "docstatus", "net_pay", "rounded_total",
			"bank_name", "bank_account_no",
		],
		order_by="employee_name asc",
	)
	details = (
		frappe.get_all(
			"Salary Detail",
			filters={"parenttype": "Salary Slip", "parent": ["in", [s.name for s in slips]]},
			fields=[
				"parent", "parentfield", "salary_component", "amount",
				"do_not_include_in_total", "statistical_component",
			],
		)
		if slips
		else []
	)
	return pe, slips, details


def _codes_by_name():
	if not frappe.get_meta("Bank").has_field("nchl_bank_code"):
		return {}
	return {
		normalize_bank_name(b.name): b.nchl_bank_code
		for b in frappe.get_all("Bank", fields=["name", "nchl_bank_code"])
		if b.nchl_bank_code
	}


def _employee_ids(employees):
	"""SSF and CIT numbers, read only from the fields the site has."""
	meta = frappe.get_meta("Employee")
	fields = [f for f in ("custom_ssf_number", "custom_cit_no", "custom_cit_number") if meta.has_field(f)]
	if not employees or not fields:
		return {}
	return {
		e.name: e
		for e in frappe.get_all(
			"Employee", filters={"name": ["in", list(employees)]}, fields=["name", *fields]
		)
	}


def _pay_amount(slip):
	rounded = not cint(frappe.db.get_single_value("Payroll Settings", "disable_rounded_total"))
	return flt(slip.rounded_total if rounded else slip.net_pay, 2)


# Latin names: these end up in bank remarks, which must be plain text.
BS_MONTHS = (
	"Baisakh", "Jestha", "Ashadh", "Shrawan", "Bhadra", "Ashwin",
	"Kartik", "Mangsir", "Poush", "Magh", "Falgun", "Chaitra",
)


def _period_label(pe):
	"""'Ashwin 2083': the BS month the period ends in."""
	try:
		import nepal_compliance.nepali_date_utils.nepali_date as nepali_date

		nepali_date._ensure_loaded()
		bs = nepali_date.ad_to_bs(getdate(pe.end_date))
		return f"{BS_MONTHS[int(bs['month']) - 1]} {bs['year']}"
	except Exception:
		return getdate(pe.end_date).strftime("%b %Y")


def _parse_list(value):
	if isinstance(value, str):
		value = frappe.parse_json(value) if value.strip().startswith("[") else [value]
	return [v for v in (value or []) if v]


# --------------------------------------------------------------------------
# Options and checks for the dialog


@frappe.whitelist()
def get_payroll_export_options(payroll_entry):
	pe, slips, details = _load(payroll_entry)
	codes = _codes_by_name()

	def totals(parentfield, hidden):
		by_root = defaultdict(float)
		for (slip, root), amount in component_amounts(details, parentfield).items():
			if root not in hidden:
				by_root[root] += amount
		return [
			{"root": root, "total": flt(total, 2)}
			for root, total in sorted(by_root.items())
			if flt(total, 2) > 0
		]

	components = totals("earnings", offset_roots(details))
	deduction_components = totals("deductions", offset_deduction_roots(details))

	no_account = [s.employee_name for s in slips if not (s.bank_account_no or "").strip()]
	no_code = sorted(
		{
			s.bank_name or _("(blank)")
			for s in slips
			if (s.bank_account_no or "").strip() and not bank_code_for(s.bank_name, codes)
		}
	)

	ids = _employee_ids({s.employee for s in slips})
	ssf = amounts_by_root(details, SSF_EMPLOYEE_ROOT)
	cit = amounts_by_root(details, CIT_ROOT)
	no_ssf_no = [
		s.employee_name for s in slips
		if ssf.get(s.name) and not (ids.get(s.employee) or {}).get("custom_ssf_number")
	]
	no_cit_no = [
		s.employee_name for s in slips
		if cit.get(s.name) and not _cit_number(ids.get(s.employee))
	]

	return {
		"period": _period_label(pe),
		"slip_count": len(slips),
		"draft_count": sum(1 for s in slips if s.docstatus == 0),
		"components": components,
		"deductions": deduction_components,
		"no_account": no_account,
		"no_bank_code": no_code,
		"no_ssf_number": no_ssf_no,
		"no_cit_number": no_cit_no,
	}


def _cit_number(employee):
	employee = employee or {}
	return employee.get("custom_cit_no") or employee.get("custom_cit_number") or ""


# --------------------------------------------------------------------------
# Download


@frappe.whitelist()
def download_payroll_export(
	payroll_entry, kind, components=None, exclude=None, deductions=None, validate_only=0
):
	"""validate_only builds the file and answers with its name, so the dialog
	can show a refusal as a message instead of a downloaded error page."""
	from frappe.desk.utils import provide_binary_file

	if kind not in KINDS:
		frappe.throw(_("Unknown download: {0}").format(kind))
	pe, slips, details = _load(payroll_entry)
	period = _period_label(pe)

	if kind in ("net_pay", "component_net_pay", "components"):
		by_component = component_amounts(details)
		chosen = set(_parse_list(exclude if kind == "net_pay" else components))
		not_paid = sorted(
			(chosen & offset_roots(details))
			| (set(_parse_list(deductions)) & offset_deduction_roots(details))
		)
		if not_paid:
			frappe.throw(
				_("{0} is deducted back and paid to the fund, not to the employee, so it cannot be paid or left out here.").format(
					", ".join(not_paid)
				)
			)
		if kind == "net_pay":
			roots = set(_parse_list(exclude))
			amounts = {
				s.name: _pay_amount(s)
				- sum(by_component.get((s.name, root), 0.0) for root in roots)
				for s in slips
			}
			label = _("Salary")
			title = _("Net Pay") + (
				" " + _("less {0}").format(", ".join(sorted(roots))) if roots else ""
			)
		elif kind == "component_net_pay":
			roots = _parse_list(components)
			taken = _parse_list(deductions)
			if not roots:
				frappe.throw(_("Tick at least one earning that makes up the payment."))
			amounts = component_net_amounts(slips, details, roots, taken)
			label = _("Salary")
			title = _("Net Pay by Component")
		else:
			roots = _parse_list(components)
			if not roots:
				frappe.throw(_("Pick at least one component to pay."))
			amounts = {
				s.name: sum(by_component.get((s.name, root), 0.0) for root in roots)
				for s in slips
			}
			label = " + ".join(roots)
			title = label
		rows, problems = bank_rows(slips, amounts, _codes_by_name(), f"{label} {period}")
		negative = [s for s, _msg in problems if flt(amounts.get(s.name)) < 0]
		if kind == "component_net_pay" and negative:
			frappe.throw(
				_("The ticked deductions are more than the ticked earnings for {0} of {1} employee(s). Tick the earnings those deductions come out of, or untick deductions.").format(
					len(negative), len(slips)
				),
				title=_("Deductions more than earnings"),
			)
		if kind == "net_pay" and negative:
			frappe.throw(
				_("Leaving out {0} takes more than net pay for {1} of {2} employee(s): those components are more than what is left after deductions. Untick the ones that are not paid separately -- usually only Festival Allowance or Bonus are.").format(
					", ".join(sorted(roots)), len(negative), len(slips)
				),
				title=_("Left out more than net pay"),
			)
		if problems:
			frappe.throw(
				_("These slips cannot go on the bank file. Fix them and download again:")
				+ "<br>"
				+ "<br>".join(
					f"{frappe.utils.escape_html(s.employee_name)}: {msg}" for s, msg in problems
				),
				title=_("Bank file incomplete"),
			)
		if not rows:
			frappe.throw(_("Nothing to pay in this payroll entry for that choice."))
		if len(rows) > BANK_MAX_ROWS:
			frappe.throw(_("The bank allows {0} rows per file; this has {1}.").format(BANK_MAX_ROWS, len(rows)))
		content = _bank_xlsx(rows)
	elif kind == "jv":
		jv = accrual_jv_view(pe)
		title = _("Accrual JV") + ("" if jv["posted"] else " " + _("Preview"))
		content = _jv_xlsx(pe, jv, period)
	elif kind == "ssf":
		title = _("SSF Contribution")
		content = _ssf_xlsx(pe, slips, details, period)
	else:
		title = _("CIT Contribution")
		content = _cit_xlsx(pe, slips, details, period)

	filename = f"{pe.name} {title} {period}"
	if cint(validate_only):
		return {"filename": f"{filename}.xlsx"}
	provide_binary_file(filename, "xlsx", content)


def _bank_xlsx(rows):
	from openpyxl import load_workbook

	wb = load_workbook(TEMPLATE)
	ws = wb.active
	for offset, row in enumerate(rows):
		r = BANK_FIRST_ROW + offset
		for col, value in enumerate(row, start=1):
			cell = ws.cell(row=r, column=col, value=value)
			# Bank code and account number are text: a leading 0 must survive.
			cell.number_format = "#,##0.00" if col == 1 else "@"
	return _save(wb)


def _contribution_xlsx(pe, period, title, headers, rows, money_from):
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

	wb = Workbook()
	ws = wb.active
	ws.title = title[:31]
	ws.append([pe.company])
	ws.append([f"{title} — {period} ({pe.start_date} to {pe.end_date}) — {pe.name}"])
	ws.append([])
	ws.append(headers)
	for row in rows:
		ws.append(row)
	totals = [_("Total")] + [""] * (money_from - 1) + [
		flt(sum(flt(r[i]) for r in rows), 2) for i in range(money_from, len(headers))
	]
	ws.append(totals)

	bold = Font(bold=True)
	thin = Side(style="thin", color="999999")
	header_fill = PatternFill("solid", fgColor="E8EEF7")
	ws["A1"].font = Font(bold=True, size=13)
	header_row, last_row = 4, ws.max_row
	for cell in ws[header_row]:
		cell.font = bold
		cell.fill = header_fill
		cell.alignment = Alignment(wrap_text=True, vertical="center")
	for row in ws.iter_rows(min_row=header_row, max_row=last_row, max_col=len(headers)):
		for cell in row:
			cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
			if cell.row > header_row and cell.column > money_from:
				cell.number_format = "#,##0.00"
			elif cell.row > header_row and cell.column > 1:
				cell.number_format = "@"
	for cell in ws[last_row]:
		cell.font = bold
	for i, header in enumerate(headers, start=1):
		ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = max(
			12, min(40, len(str(header)) + 4)
		)
	ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
	return _save(wb)


def _ssf_xlsx(pe, slips, details, period):
	employee_share = amounts_by_root(details, SSF_EMPLOYEE_ROOT)
	employer_share = amounts_by_root(details, SSF_EMPLOYER_ROOT)
	ssf_base = _slip_values(slips, "custom_ssf_base")
	ids = _employee_ids({s.employee for s in slips})
	rows = []
	for s in slips:
		employee, employer = employee_share.get(s.name, 0.0), employer_share.get(s.name, 0.0)
		if not (employee or employer):
			continue
		rows.append(
			[
				len(rows) + 1,
				s.employee,
				s.employee_name,
				(ids.get(s.employee) or {}).get("custom_ssf_number") or "",
				flt(ssf_base.get(s.name), 2),
				employee,
				employer,
				flt(employee + employer, 2),
			]
		)
	headers = [
		_("S.N."), _("Employee ID"), _("Employee Name"), _("SSF Number"), _("SSF Base"),
		_("Employee Contribution"), _("Employer Contribution"), _("Total Contribution"),
	]
	return _contribution_xlsx(pe, period, _("SSF Contribution"), headers, rows, money_from=4)


def _cit_xlsx(pe, slips, details, period):
	cit = amounts_by_root(details, CIT_ROOT)
	ids = _employee_ids({s.employee for s in slips})
	rows = []
	for s in slips:
		amount = cit.get(s.name, 0.0)
		if not amount:
			continue
		rows.append(
			[len(rows) + 1, s.employee, s.employee_name, _cit_number(ids.get(s.employee)), amount]
		)
	headers = [_("S.N."), _("Employee ID"), _("Employee Name"), _("CIT Number"), _("CIT Amount")]
	return _contribution_xlsx(pe, period, _("CIT Contribution"), headers, rows, money_from=4)


def _slip_values(slips, field):
	if not slips or not frappe.get_meta("Salary Slip").has_field(field):
		return {}
	return dict(
		frappe.get_all(
			"Salary Slip",
			filters={"name": ["in", [s.name for s in slips]]},
			fields=["name", field],
			as_list=True,
		)
	)


def _save(wb):
	out = BytesIO()
	wb.save(out)
	return out.getvalue()


# --------------------------------------------------------------------------
# Accrual Journal Entry view


JV_FIELDS = (
	("account", "Account"),
	("cost_center", "Cost Center"),
	("party_type", "Party Type"),
	("party", "Party"),
	("reference_type", "Reference Type"),
	("reference_name", "Reference"),
)


def accrual_jv_view(pe):
	"""The accrual Journal Entry's rows: the posted one, or what would be posted.

	A posted JV is read back as it stands. Otherwise HRMS's own make_accrual_jv_entry
	runs over every slip of this run not yet in a JV -- drafts included, since
	the JV is posted the moment they are submitted -- and its rows are caught at
	make_journal_entry, before anything is saved. The whole run is rolled back
	regardless.
	"""
	from nepal_hrms.custom_code.payroll.accrual_jv import get_accrual_journal_entries

	posted = next((e for e in get_accrual_journal_entries(pe.name) if e.docstatus == 1), None)
	if posted:
		je = frappe.get_doc("Journal Entry", posted.name)
		return {
			"posted": je.name,
			"posting_date": je.posting_date,
			"remark": je.user_remark,
			"accounts": [a.as_dict() for a in je.accounts],
		}

	captured = {}

	def unaccrued_slips(ss_status, as_dict=False):
		# Drafts and submitted slips alike; the stock query skips any slip a JV holds.
		return type(pe).get_sal_slip_list(pe, 0, as_dict) + type(pe).get_sal_slip_list(
			pe, 1, as_dict
		)

	def capture(accounts, currencies, payroll_payable_account=None, **kwargs):
		captured["accounts"] = [frappe._dict(a) for a in accounts]
		captured["remark"] = kwargs.get("user_remark")

	frappe.db.savepoint("nh_jv_preview")
	try:
		# Instance attributes shadow the methods for this one run; reading the
		# preview needs read access only, which _load has already checked.
		pe.get_sal_slip_list = unaccrued_slips
		pe.make_journal_entry = capture
		pe.check_permission = lambda *args, **kwargs: None
		pe.make_accrual_jv_entry(submitted_salary_slips=None)
	finally:
		for attr in ("get_sal_slip_list", "make_journal_entry", "check_permission"):
			pe.__dict__.pop(attr, None)
		frappe.db.rollback(save_point="nh_jv_preview")

	if not captured.get("accounts"):
		frappe.throw(_("Every salary slip of this Payroll Entry is already in a Journal Entry, and none of them is a posted accrual entry to show."))
	return {
		"posted": None,
		"posting_date": pe.posting_date,
		"remark": captured.get("remark"),
		"accounts": captured["accounts"],
	}


def _jv_xlsx(pe, jv, period):
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

	accounts = sorted(
		jv["accounts"],
		# Debits (expenses) first, then credits; within each, by account.
		key=lambda a: (
			0 if flt(a.get("debit_in_account_currency")) else 1,
			a.get("account") or "",
			a.get("party") or "",
		),
	)
	debit_total = flt(sum(flt(a.get("debit_in_account_currency")) for a in accounts), 2)
	credit_total = flt(sum(flt(a.get("credit_in_account_currency")) for a in accounts), 2)

	wb = Workbook()
	ws = wb.active
	ws.title = "Journal Entry"
	status = (
		_("Posted as {0}").format(jv["posted"])
		if jv["posted"]
		else _("PREVIEW -- not posted. It is posted when the salary slips are submitted.")
	)
	ws.append([_("Accrual Journal Entry"), pe.company])
	ws.append([_("Status"), status])
	ws.append([_("Payroll Entry"), f"{pe.name} -- {period} ({pe.start_date} to {pe.end_date})"])
	ws.append([_("Posting Date"), str(jv["posting_date"] or "")])
	ws.append([_("Remark"), jv["remark"] or ""])
	ws.append([])
	header_row = ws.max_row + 1
	ws.append([_("#")] + [_(label) for _f, label in JV_FIELDS] + [_("Debit"), _("Credit")])
	for i, a in enumerate(accounts, start=1):
		ws.append(
			[i]
			+ [a.get(field) or "" for field, _label in JV_FIELDS]
			+ [
				flt(a.get("debit_in_account_currency"), 2) or None,
				flt(a.get("credit_in_account_currency"), 2) or None,
			]
		)
	ws.append([""] * (len(JV_FIELDS) + 1) + [debit_total, credit_total])
	total_row = ws.max_row
	ws.cell(row=total_row, column=2, value=_("Total"))
	if abs(debit_total - credit_total) > 0.005:
		ws.append([""] * (len(JV_FIELDS)) + [_("Out of balance by"), flt(debit_total - credit_total, 2)])

	# Expense summary: what this payroll costs, by account.
	by_account = defaultdict(float)
	for a in accounts:
		if flt(a.get("debit_in_account_currency")):
			by_account[a.get("account")] += flt(a.get("debit_in_account_currency"))
	summary = wb.create_sheet(_("Debits by Account")[:31])
	summary.append([_("Account"), _("Debit")])
	for account, amount in sorted(by_account.items()):
		summary.append([account, flt(amount, 2)])
	summary.append([_("Total"), flt(sum(by_account.values()), 2)])

	bold = Font(bold=True)
	thin = Side(style="thin", color="999999")
	fill = PatternFill("solid", fgColor="E8EEF7")
	ws["A1"].font = Font(bold=True, size=13)
	ws["B2"].font = Font(bold=True, color="C0392B" if not jv["posted"] else "1E8449")
	for col in range(1, 4):
		for row in range(1, 6):
			ws.cell(row=row, column=col).alignment = Alignment(vertical="top")
	for sheet, first, last, money_from in (
		(ws, header_row, total_row, len(JV_FIELDS) + 2),
		(summary, 1, summary.max_row, 2),
	):
		for cell in sheet[first]:
			cell.font = bold
			cell.fill = fill
		for row in sheet.iter_rows(min_row=first, max_row=last):
			for cell in row:
				cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
				if cell.row > first and cell.column >= money_from:
					cell.number_format = "#,##0.00"
		for cell in sheet[last]:
			cell.font = bold
	widths = {"A": 6, "B": 42, "C": 26, "D": 12, "E": 18, "F": 16, "G": 22, "H": 16, "I": 16}
	for col, width in widths.items():
		ws.column_dimensions[col].width = width
	summary.column_dimensions["A"].width = 48
	summary.column_dimensions["B"].width = 18
	ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
	return _save(wb)
