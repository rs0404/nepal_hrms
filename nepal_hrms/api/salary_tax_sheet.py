# Copyright (c) 2026, Yarsa Labs Pvt. Ltd. and contributors
# For license information, please see license.txt

"""Salary & Tax Sheet: every slip of a Payroll Entry, laid out so the arithmetic
can be followed from earnings to net pay and from annual income to this month's TDS.

Sheets
- Salary Sheet: one row per employee -- attendance, each earning and deduction,
  net pay, the annual tax working, and TDS month by month.
- Tax Calculation: each employee's tax worked step by step, with the slab table.
- TDS by Month: the fiscal year's twelve months; deducted, this month, projected.
- How to read: what the colours mean.

Every figure is the slip's own. The slab table is re-worked from the employee's
Income Tax Slab the way HRMS does it, as an explanation; where the slip's annual
tax differs (marginal relief, conditional slabs) the sheet says so instead of
silently showing two numbers.
"""

from collections import defaultdict

import frappe
from frappe import _
from frappe.utils import cint, cstr, flt, getdate

from nepal_hrms.api.payroll_exports import (
	BS_MONTHS,
	NP_MONEY_FORMAT,
	_save,
	component_root,
	counted_in_totals,
)
from nepal_hrms.custom_code.payroll.component_variant import structure_variant

SLIP_FIELDS = (
	"name", "employee", "employee_name", "department", "designation", "salary_structure",
	"start_date", "end_date", "docstatus",
	"total_working_days", "payment_days", "absent_days", "custom_penalty_days", "leave_without_pay",
	"gross_pay", "total_deduction", "net_pay",
	"total_earnings", "income_from_other_sources", "non_taxable_earnings",
	"standard_tax_exemption_amount", "gross_year_to_date",
	"custom_assessable_employment_income", "custom_total_retirement_contribution",
	"custom_one_third_assessable_income", "custom_annual_retirement_limit",
	"custom_allowed_retirement_deduction", "custom_annual_taxable_deductions",
	"annual_taxable_amount", "custom_projected_annual_tds", "total_income_tax",
	"custom_tds_paid_till_date", "income_tax_deducted_till_date",
	"custom_remaining_payroll_periods", "custom_current_period_tds", "current_month_income_tax",
)

# (header fill, light fill)
COLOURS = {
	"emp": ("5D6D7E", "F2F4F4"),
	"att": ("7F8C8D", "F4F6F6"),
	"earn": ("1E8449", "E9F7EF"),
	"ded": ("B03A2E", "FDEDEC"),
	"net": ("1F4E79", "D6EAF8"),
	"tax": ("6C3483", "F4ECF7"),
	"tds": ("117A65", "E8F8F5"),
	"info": ("A6ACAF", "F8F9F9"),
}


# --------------------------------------------------------------------------
# Pure helpers (unit-tested without a site)


def ordered_roots(details, parentfield, counted=True):
	"""Component roots of one table in slip order: first by the lowest row they
	ever sit on, then by name, so Basic and Grade lead as they do on a slip."""
	first_idx = {}
	for d in details:
		if d.parentfield != parentfield or counted_in_totals(d) != counted:
			continue
		root = component_root(d.salary_component)
		first_idx[root] = min(first_idx.get(root, 10**6), cint(d.get("idx")) or 10**6)
	return sorted(first_idx, key=lambda root: (first_idx[root], root))


def amounts_by_slip(details, parentfield, counted=True):
	"""{(slip, root): amount} for counted rows, or for Calculation Only rows."""
	totals = defaultdict(float)
	for d in details:
		if d.parentfield == parentfield and counted_in_totals(d) == counted:
			totals[(d.parent, component_root(d.salary_component))] += flt(d.amount)
	return {key: flt(value, 2) for key, value in totals.items()}


def slab_breakdown(income, slabs):
	"""[(from, to, rate, income in band, tax)] and the total, as HRMS works it.

	HRMS bands are inclusive at both ends (to - from + 1). A slab with a
	condition is listed but not charged: whether it applies depends on the slip.
	"""
	rows, total = [], 0.0
	income = flt(income)
	for s in slabs:
		low, high, rate = flt(s.from_amount), flt(s.to_amount), flt(s.percent_deduction)
		band = 0.0
		if not cstr(s.get("condition")).strip() and income >= low:
			band = (income - low + 1) if (not high or income < high) else (high - low + 1)
		tax = flt(band * rate / 100, 2)
		rows.append((low, high, rate, flt(band, 2), tax))
		total += tax
	return rows, flt(total, 2)


def np_number(value):
	"""12,34,567 -- Nepali grouping, for figures written into text."""
	whole = str(int(round(abs(flt(value)))))
	head, tail = whole[:-3], whole[-3:]
	groups = []
	while len(head) > 2:
		groups.insert(0, head[-2:])
		head = head[:-2]
	if head:
		groups.insert(0, head)
	return ("-" if flt(value) < 0 else "") + ",".join(groups + [tail])


def fiscal_months(period_start):
	"""The twelve BS month names of a fiscal year, starting at its first month."""
	first = _bs(period_start)["month"]
	return [BS_MONTHS[(first - 1 + i) % 12] for i in range(12)]


def _bs(ad_date):
	import nepal_compliance.nepali_date_utils.nepali_date as nepali_date

	nepali_date._ensure_loaded()
	return nepali_date.ad_to_bs(getdate(ad_date))


def bs_month_name(ad_date):
	return BS_MONTHS[int(_bs(ad_date)["month"]) - 1]


# --------------------------------------------------------------------------
# Data


def _load(pe):
	columns = set(frappe.db.get_table_columns("Salary Slip"))
	fields = [f for f in SLIP_FIELDS if f in columns]
	slips = frappe.get_all(
		"Salary Slip",
		filters={"payroll_entry": pe.name, "docstatus": ["<", 2]},
		fields=fields,
		order_by="employee_name asc",
	)
	for s in slips:
		for f in SLIP_FIELDS:
			s.setdefault(f, None)
	details = (
		frappe.get_all(
			"Salary Detail",
			filters={"parenttype": "Salary Slip", "parent": ["in", [s.name for s in slips]]},
			fields=[
				"parent", "parentfield", "salary_component", "amount", "idx",
				"do_not_include_in_total", "statistical_component",
			],
		)
		if slips
		else []
	)
	return slips, details


def _payroll_period(pe):
	return (
		frappe.get_all(
			"Payroll Period",
			filters={"company": pe.company, "start_date": ["<=", pe.end_date], "end_date": [">=", pe.end_date]},
			fields=["name", "start_date", "end_date"],
			order_by="start_date desc",
			limit=1,
		)
		or [None]
	)[0]


def _tax_components():
	return set(frappe.get_all("Salary Component", {"variable_based_on_taxable_salary": 1}, pluck="name"))


def _earlier_tds(slips, period, tax_components):
	"""{employee: {BS month: TDS}} from submitted slips earlier in the fiscal year."""
	if not period or not slips or not tax_components:
		return {}
	earliest = min(getdate(s.start_date) for s in slips)
	rows = frappe.db.sql(
		"""
		select ss.employee, ss.end_date, sum(sd.amount) as tds
		from `tabSalary Slip` ss
		join `tabSalary Detail` sd on sd.parent = ss.name and sd.parenttype = 'Salary Slip'
		where ss.docstatus = 1
			and ss.employee in %(employees)s
			and ss.start_date >= %(period_start)s
			and ss.end_date < %(before)s
			and sd.parentfield = 'deductions'
			and sd.salary_component in %(components)s
		group by ss.name
		""",
		{
			"employees": tuple({s.employee for s in slips}),
			"period_start": period.start_date,
			"before": earliest,
			"components": tuple(tax_components),
		},
		as_dict=True,
	)
	out = defaultdict(lambda: defaultdict(float))
	for r in rows:
		out[r.employee][bs_month_name(r.end_date)] += flt(r.tds)
	return out


def _slabs_by_employee(slips):
	"""{employee: (slab name, slab rows)} from the assignment in force for the slip."""
	out, cache = {}, {}
	for s in slips:
		slab = frappe.db.get_value(
			"Salary Structure Assignment",
			{"employee": s.employee, "docstatus": 1, "from_date": ["<=", s.end_date]},
			"income_tax_slab",
			order_by="from_date desc",
		)
		if slab and slab not in cache:
			cache[slab] = frappe.get_all(
				"Taxable Salary Slab",
				filters={"parent": slab, "parenttype": "Income Tax Slab"},
				fields=["from_amount", "to_amount", "percent_deduction", "condition"],
				order_by="idx asc",
			)
		out[s.employee] = (slab, cache.get(slab) or [])
	return out


def _tax(s, tds_this_month):
	"""The slip's annual tax working, falling back to stock HRMS fields."""
	assessable = flt(s.custom_assessable_employment_income) or (
		flt(s.total_earnings) + flt(s.income_from_other_sources)
		- flt(s.non_taxable_earnings) - flt(s.standard_tax_exemption_amount)
	)
	allowed = flt(s.custom_allowed_retirement_deduction)
	deductions = flt(s.custom_annual_taxable_deductions)
	return frappe._dict(
		projected_earnings=flt(s.total_earnings) + flt(s.income_from_other_sources),
		non_taxable=flt(s.non_taxable_earnings),
		exemption=flt(s.standard_tax_exemption_amount),
		assessable=flt(assessable, 2),
		retirement=flt(s.custom_total_retirement_contribution),
		one_third=flt(s.custom_one_third_assessable_income),
		cap=flt(s.custom_annual_retirement_limit),
		allowed=allowed,
		other_deductions=flt(max(deductions - allowed, 0), 2),
		taxable=flt(s.annual_taxable_amount, 2),
		annual_tax=flt(s.custom_projected_annual_tds or s.total_income_tax, 2),
		paid=flt(s.custom_tds_paid_till_date or s.income_tax_deducted_till_date, 2),
		remaining=cint(s.custom_remaining_payroll_periods),
		this_month=flt(tds_this_month, 2),
	)


# --------------------------------------------------------------------------
# Workbook


def build_salary_tax_sheet(pe, period_label):
	from openpyxl import Workbook

	slips, details = _load(pe)
	if not slips:
		frappe.throw(_("This Payroll Entry has no salary slips yet."))

	earnings = ordered_roots(details, "earnings")
	deductions = ordered_roots(details, "deductions")
	info = ordered_roots(details, "deductions", counted=False) + ordered_roots(
		details, "earnings", counted=False
	)
	earn_amt = amounts_by_slip(details, "earnings")
	ded_amt = amounts_by_slip(details, "deductions")
	info_amt = {
		**amounts_by_slip(details, "earnings", counted=False),
		**amounts_by_slip(details, "deductions", counted=False),
	}

	tax_components = _tax_components()
	tax_roots = {component_root(c) for c in tax_components}
	period = _payroll_period(pe)
	months = fiscal_months(period.start_date) if period else []
	current_month = bs_month_name(pe.end_date)
	earlier = _earlier_tds(slips, period, tax_components)
	slabs = _slabs_by_employee(slips)

	rows = []
	for s in slips:
		tds_now = sum(ded_amt.get((s.name, root), 0.0) for root in tax_roots)
		rows.append(
			frappe._dict(
				slip=s,
				type=structure_variant(s.salary_structure) or "",
				tax=_tax(s, tds_now),
				earlier=earlier.get(s.employee, {}),
				slab=slabs.get(s.employee, (None, [])),
			)
		)

	ctx = frappe._dict(
		pe=pe, period_label=period_label, months=months, current_month=current_month,
		earnings=earnings, deductions=deductions, info=info,
		earn_amt=earn_amt, ded_amt=ded_amt, info_amt=info_amt, rows=rows,
	)
	wb = Workbook()
	_salary_sheet(wb.active, ctx)
	_tax_sheet(wb.create_sheet("Tax Calculation"), ctx)
	_tds_sheet(wb.create_sheet("TDS by Month"), ctx)
	_legend_sheet(wb.create_sheet("How to read"), ctx)
	return _save(wb)


def _styles():
	from openpyxl.styles import Border, Font, PatternFill, Side

	thin = Side(style="thin", color="BFC9CA")
	return frappe._dict(
		border=Border(top=thin, bottom=thin, left=thin, right=thin),
		white_bold=Font(bold=True, color="FFFFFF"),
		bold=Font(bold=True),
		muted=Font(italic=True, color="7F8C8D"),
		fill=lambda hex_: PatternFill("solid", fgColor=hex_),
		Font=Font,
	)


def _title(ws, ctx, title):
	from openpyxl.styles import Font

	pe = ctx.pe
	ws["A1"] = title
	ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
	ws["A2"] = _("{0}  ·  Payroll Entry {1}  ·  {2} ({3} to {4})").format(
		pe.company, pe.name, ctx.period_label, pe.start_date, pe.end_date
	)
	ws["A2"].font = Font(italic=True, color="555555")
	drafts = sum(1 for r in ctx.rows if r.slip.docstatus == 0)
	if drafts:
		ws["A3"] = _("{0} of {1} salary slip(s) are still drafts; their figures may change.").format(
			drafts, len(ctx.rows)
		)
		ws["A3"].font = Font(bold=True, color="B03A2E")


def _salary_sheet(ws, ctx):
	from openpyxl.styles import Alignment
	from openpyxl.utils import get_column_letter

	st = _styles()
	ws.title = _("Salary Sheet")
	_title(ws, ctx, _("Salary & Tax Sheet"))
	earlier_months = ctx.months[: ctx.months.index(ctx.current_month)] if ctx.current_month in ctx.months else []

	# (group, colour key, [(label, getter, kind)]); kind: text | days | money
	def slip(field):
		return lambda r: r.slip.get(field)

	def tax(field):
		return lambda r: r.tax.get(field)

	groups = [
		(_("Employee"), "emp", [
			(_("Employee ID"), slip("employee"), "text"),
			(_("Employee Name"), slip("employee_name"), "text"),
			(_("Department"), slip("department"), "text"),
			(_("Designation"), slip("designation"), "text"),
			(_("Type"), lambda r: r.type, "text"),
		]),
		(_("Attendance"), "att", [
			(_("Working Days"), slip("total_working_days"), "days"),
			(_("Payment Days"), slip("payment_days"), "days"),
			(_("Absent Days"), slip("absent_days"), "days"),
			(_("Penalty Days"), slip("custom_penalty_days"), "days"),
			(_("LWP"), slip("leave_without_pay"), "days"),
		]),
		(_("Earnings"), "earn", [
			(root, (lambda root: lambda r: ctx.earn_amt.get((r.slip.name, root), 0.0))(root), "money")
			for root in ctx.earnings
		] + [(_("Gross Earnings"), slip("gross_pay"), "money")]),
		(_("Deductions"), "ded", [
			(root, (lambda root: lambda r: ctx.ded_amt.get((r.slip.name, root), 0.0))(root), "money")
			for root in ctx.deductions
		] + [(_("Total Deductions"), slip("total_deduction"), "money")]),
		(_("Net Pay"), "net", [(_("Net Pay"), slip("net_pay"), "money")]),
		(_("Tax Calculation (annual)"), "tax", [
			(_("Assessable Income"), tax("assessable"), "money"),
			(_("Retirement Contribution"), tax("retirement"), "money"),
			(_("1/3 of Income"), tax("one_third"), "money"),
			(_("Annual Cap"), tax("cap"), "cap"),
			(_("Allowed Deduction"), tax("allowed"), "money"),
			(_("Taxable Income"), tax("taxable"), "money"),
			(_("Annual Tax"), tax("annual_tax"), "money"),
			(_("TDS Paid Till Last Month"), tax("paid"), "money"),
			(_("Remaining Months"), tax("remaining"), "days"),
		]),
		(_("TDS by Month"), "tds", [
			(month, (lambda month: lambda r: flt(r.earlier.get(month), 2))(month), "money")
			for month in earlier_months
		] + [
			(_("{0} (this)").format(ctx.current_month), tax("this_month"), "money"),
			(_("TDS This Year"), lambda r: flt(sum(r.earlier.values()) + r.tax.this_month, 2), "money"),
		]),
	]
	if ctx.info:
		groups.append((_("Calculation Only"), "info", [
			(root, (lambda root: lambda r: ctx.info_amt.get((r.slip.name, root), 0.0))(root), "info")
			for root in ctx.info
		]))

	group_row, head_row = 5, 6
	columns, col = [], 1
	for c, label in ((1, _("S.N.")),):
		cell = ws.cell(row=head_row, column=c, value=label)
		cell.font, cell.fill, cell.border = st.white_bold, st.fill(COLOURS["emp"][0]), st.border
		columns.append((c, "emp", None, "sn"))
	col = 2
	for gname, key, cols in groups:
		start = col
		for label, getter, kind in cols:
			cell = ws.cell(row=head_row, column=col, value=label)
			cell.font, cell.fill, cell.border = st.white_bold, st.fill(COLOURS[key][0]), st.border
			cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
			columns.append((col, key, getter, kind))
			col += 1
		ws.merge_cells(start_row=group_row, start_column=start, end_row=group_row, end_column=col - 1)
		g = ws.cell(row=group_row, column=start, value=gname)
		g.font, g.fill = st.white_bold, st.fill(COLOURS[key][0])
		g.alignment = Alignment(horizontal="center")
	ws.cell(row=group_row, column=1).fill = st.fill(COLOURS["emp"][0])

	key_figures = {_("Gross Earnings"), _("Total Deductions"), _("Net Pay"), _("Annual Tax")}
	for i, r in enumerate(ctx.rows, start=1):
		row = head_row + i
		for c, key, getter, kind in columns:
			value = i if kind == "sn" else getter(r)
			if kind in ("money", "cap", "info"):
				value = flt(value, 2)
			elif kind == "days":
				value = flt(value)
			cell = ws.cell(row=row, column=c, value=value)
			cell.border, cell.fill = st.border, st.fill(COLOURS[key][1])
			if kind in ("money", "cap", "info"):
				cell.number_format = NP_MONEY_FORMAT
			elif kind == "days":
				cell.number_format = "General"  # 22, 22.5 -- "0.##" shows a whole number as "22."
			header = ws.cell(row=head_row, column=c).value
			if header in key_figures or (key == "tds" and header.endswith(_("(this)"))):
				cell.font = st.Font(bold=True, color="1F4E79" if header == _("Net Pay") else "000000")
			elif kind == "info":
				cell.font = st.muted

	total_row = head_row + len(ctx.rows) + 1
	ws.cell(row=total_row, column=3, value=_("Total"))
	for c, key, getter, kind in columns:
		cell = ws.cell(row=total_row, column=c)
		if kind in ("money", "info"):
			cell.value = flt(sum(flt(ws.cell(row=r, column=c).value) for r in range(head_row + 1, total_row)), 2)
			cell.number_format = NP_MONEY_FORMAT
		cell.font, cell.fill, cell.border = st.bold, st.fill("D5D8DC"), st.border

	widths = {1: 5, 2: 12, 3: 24, 4: 14, 5: 18, 6: 8}
	for c in range(1, col):
		ws.column_dimensions[get_column_letter(c)].width = widths.get(c, 13)
	ws.row_dimensions[head_row].height = 48
	ws.freeze_panes = ws.cell(row=head_row + 1, column=4)


def _tax_sheet(ws, ctx):
	from openpyxl.styles import Font

	st = _styles()
	_title(ws, ctx, _("How each employee's TDS is worked out"))
	r = 5
	for row in ctx.rows:
		s, t = row.slip, row.tax
		slab_name, slab_rows = row.slab
		ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
		head = ws.cell(row=r, column=1, value=f"{s.employee} · {s.employee_name} · {row.type or ''}".rstrip(" ·"))
		head.font, head.fill = st.white_bold, st.fill(COLOURS["tax"][0])
		r += 1

		lines = [
			(_("Income for the year"), None, None, False),
			(_("Projected annual earnings"), t.projected_earnings, _("earned so far + this month + the rest of the year at this month's rate"), False),
			(_("Less: non-taxable earnings"), t.non_taxable, "", False),
		]
		if t.exemption:
			lines.append((_("Less: standard exemption"), t.exemption, "", False))
		lines += [
			(_("Assessable income"), t.assessable, _("projected earnings − non-taxable"), True),
			(_("Retirement deduction"), None, None, False),
			(_("SSF + CIT for the year"), t.retirement, "", False),
			(_("One-third of assessable income"), t.one_third, _("assessable ÷ 3"), False),
			(_("Annual cap"), t.cap, "", False),
			(_("Allowed deduction"), t.allowed, _("smallest of the three"), True),
		]
		if t.other_deductions:
			lines.append((_("Other taxable deductions"), t.other_deductions, _("insurance and other allowed deductions"), False))
		lines.append((_("Taxable income"), t.taxable, _("assessable − deductions"), True))

		for label, value, how, strong in lines:
			a = ws.cell(row=r, column=1, value=label)
			if value is None:
				a.font = Font(bold=True, color=COLOURS["tax"][0])
			else:
				v = ws.cell(row=r, column=2, value=flt(value, 2))
				v.number_format = NP_MONEY_FORMAT
				ws.cell(row=r, column=3, value=how).font = st.muted
				if strong:
					a.font = v.font = st.bold
					for c in (1, 2, 3):
						ws.cell(row=r, column=c).fill = st.fill(COLOURS["tax"][1])
			r += 1

		ws.cell(row=r, column=1, value=_("Tax slab: {0}").format(slab_name or _("none found"))).font = st.muted
		r += 1
		for c, label in enumerate([_("Slab"), _("Rate"), _("Income in slab"), _("Tax")], start=1):
			cell = ws.cell(row=r, column=c, value=label)
			cell.font, cell.fill = st.white_bold, st.fill("8E44AD")
		r += 1
		breakdown, slab_tax = slab_breakdown(t.taxable, slab_rows)
		for low, high, rate, band, amount in breakdown:
			ws.cell(row=r, column=1, value=f"{np_number(low)} – {np_number(high)}" if high else _("above {0}").format(np_number(low)))
			ws.cell(row=r, column=2, value=rate / 100).number_format = "0%"
			ws.cell(row=r, column=3, value=band).number_format = NP_MONEY_FORMAT
			ws.cell(row=r, column=4, value=amount).number_format = NP_MONEY_FORMAT
			if not band:
				for c in range(1, 5):
					ws.cell(row=r, column=c).font = Font(color="AAB7B8")
			r += 1

		closing = [
			(_("Annual tax"), t.annual_tax, _("as calculated on the slip")),
			(_("TDS paid till last month"), t.paid, ""),
			(_("Months left including this one"), t.remaining, ""),
			(_("TDS this month"), t.this_month, _("(annual tax − paid) ÷ months left")),
		]
		for label, value, how in closing:
			a = ws.cell(row=r, column=1, value=label)
			v = ws.cell(row=r, column=2, value=value)
			v.number_format = "0" if label == _("Months left including this one") else NP_MONEY_FORMAT
			ws.cell(row=r, column=3, value=how).font = st.muted
			if label == _("TDS this month"):
				for c in (1, 2, 3):
					ws.cell(row=r, column=c).fill = st.fill(COLOURS["net"][1])
				a.font = v.font = Font(bold=True, color="1F4E79")
			r += 1
		if abs(slab_tax - t.annual_tax) > 1:
			ws.cell(
				row=r, column=1,
				value=_("The slab table gives {0}; the slip's annual tax is {1} (marginal relief, a conditional slab or other taxes).").format(
					np_number(slab_tax), np_number(t.annual_tax)
				),
			).font = Font(italic=True, color="B03A2E")
			r += 1
		r += 1

	for col, width in (("A", 46), ("B", 18), ("C", 46), ("D", 16)):
		ws.column_dimensions[col].width = width


def _tds_sheet(ws, ctx):
	from openpyxl.styles import Alignment, Font
	from openpyxl.utils import get_column_letter

	st = _styles()
	_title(ws, ctx, _("TDS month by month"))
	months = ctx.months or [ctx.current_month]
	heads = [_("Employee ID"), _("Employee Name")] + months + [
		_("Deducted so far"), _("Projected rest of year"), _("Annual tax"),
	]
	for c, label in enumerate(heads, start=1):
		cell = ws.cell(row=5, column=c, value=label)
		cell.font, cell.fill = st.white_bold, st.fill(COLOURS["tds"][0])
		cell.alignment = Alignment(horizontal="center", wrap_text=True)
	current = months.index(ctx.current_month) if ctx.current_month in months else len(months) - 1
	for i, row in enumerate(ctx.rows, start=6):
		t = row.tax
		ws.cell(row=i, column=1, value=row.slip.employee)
		ws.cell(row=i, column=2, value=row.slip.employee_name)
		done = flt(sum(row.earlier.values()) + t.this_month, 2)
		rest = flt(max(t.annual_tax - done, 0), 2)
		later = max(t.remaining - 1, 0)
		for j, month in enumerate(months):
			cell = ws.cell(row=i, column=3 + j)
			if j < current:
				cell.value = flt(row.earlier.get(month), 2)
				cell.fill = st.fill(COLOURS["tds"][1])
			elif j == current:
				cell.value = t.this_month
				cell.fill = st.fill("FCF3CF")
				cell.font = st.bold
			elif j - current <= later:
				cell.value = flt(rest / later, 2) if later else 0.0
				cell.font = Font(italic=True, color="95A5A6")
			cell.number_format = NP_MONEY_FORMAT
		for k, value in enumerate([done, rest, t.annual_tax]):
			cell = ws.cell(row=i, column=3 + len(months) + k, value=value)
			cell.number_format, cell.font = NP_MONEY_FORMAT, st.bold
	ws.column_dimensions["A"].width = 12
	ws.column_dimensions["B"].width = 24
	for c in range(3, len(heads) + 1):
		ws.column_dimensions[get_column_letter(c)].width = 12
	ws.freeze_panes = "C6"


def _legend_sheet(ws, ctx):
	st = _styles()
	_title(ws, ctx, _("How to read this file"))
	legend = [
		("emp", _("Employee details")),
		("att", _("Attendance the slip used")),
		("earn", _("Earnings paid this month; Gross Earnings is the slip's gross")),
		("ded", _("Deductions taken this month; Total Deductions is the slip's total")),
		("net", _("Net Pay = Gross Earnings − Total Deductions")),
		("tax", _("Annual tax working, as on the slip's TDS breakdown")),
		("tds", _("TDS deducted in earlier months of the fiscal year, and this month")),
		("info", _("Calculation Only: already taken out of earnings, not deducted again")),
	]
	for i, (key, text) in enumerate(legend, start=5):
		ws.cell(row=i, column=1).fill = st.fill(COLOURS[key][0])
		ws.cell(row=i, column=2, value=text)
	ws.cell(
		row=5 + len(legend) + 1, column=2,
		value=_("TDS by Month: green = already deducted, yellow = this month, grey italic = the rest of the annual tax spread over the months left."),
	)
	ws.column_dimensions["A"].width = 4
	ws.column_dimensions["B"].width = 100
