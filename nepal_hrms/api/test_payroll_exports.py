from io import BytesIO
from unittest import TestCase

from frappe import _dict
from openpyxl import load_workbook

from nepal_hrms.api.payroll_exports import (
	BANK_FIRST_ROW,
	NP_MONEY_FORMAT,
	_bank_xlsx,
	amounts_by_root,
	bank_code_for,
	bank_rows,
	component_amounts,
	component_net_amounts,
	component_root,
	fill_missing_accounts,
	normalize_bank_name,
	offset_deduction_roots,
	offset_roots,
)


def detail(parent, parentfield, component, amount):
	return _dict(parent=parent, parentfield=parentfield, salary_component=component, amount=amount)


def slip(name, account="0010017501", bank="Nabil Bank", employee="EMP-1", employee_name="A Person"):
	return _dict(
		name=name, bank_account_no=account, bank_name=bank,
		employee=employee, employee_name=employee_name,
	)


CODES = {normalize_bank_name("Nabil Bank Limited"): "0401"}


class TestComponentGrouping(TestCase):
	def test_variants_share_a_root(self):
		self.assertEqual(component_root("Festival Allowance - Wages"), "Festival Allowance")
		self.assertEqual(component_root("Festival Allowance - Salary"), "Festival Allowance")
		self.assertEqual(component_root("CIT"), "CIT")
		self.assertEqual(
			component_root("Employer's Contribution SSF Deduction"),
			"Employer's Contribution SSF Deduction",
		)

	def test_component_amounts_sum_variants_per_slip_and_round_once(self):
		amounts = component_amounts(
			[
				detail("S1", "earnings", "Festival Allowance - Salary", 100.004),
				detail("S1", "earnings", "Festival Allowance - Wages", 0.004),
				detail("S1", "deductions", "CIT", 50),
			]
		)
		self.assertEqual(amounts, {("S1", "Festival Allowance"): 100.01})

	def test_employer_ssf_counts_the_earning_not_its_offsetting_deduction(self):
		details = [
			detail("S1", "earnings", "Employer's Contribution SSF - Salary", 2000),
			detail("S1", "deductions", "Employer's Contribution SSF Deduction", 2000),
		]
		self.assertEqual(amounts_by_root(details, "Employer's Contribution SSF"), {"S1": 2000})


class TestOffsetRoots(TestCase):
	def test_employer_contribution_deducted_back_is_offset(self):
		details = [
			detail("S1", "earnings", "Employer's Contribution SSF - Wages", 2000),
			detail("S1", "deductions", "Employer's Contribution SSF Deduction", 2000),
			detail("S1", "earnings", "Festival Allowance - Wages", 5000),
			detail("S1", "deductions", "CIT", 500),
		]
		self.assertEqual(offset_roots(details), {"Employer's Contribution SSF"})


class TestComponentNetPay(TestCase):
	details = [
		detail("S1", "earnings", "Basic Salary - Salary", 40000),
		detail("S1", "earnings", "Festival Allowance - Salary", 30000),
		detail("S1", "earnings", "Employer's Contribution SSF - Salary", 8000),
		detail("S1", "deductions", "Employer's Contribution SSF Deduction", 8000),
		detail("S1", "deductions", "Employee's Contribution SSF", 4400),
		_dict(
			parent="S1", parentfield="deductions", salary_component="Leave and Late Deduction - Salary",
			amount=1500, do_not_include_in_total=1,
		),
	]
	slips = [_dict(name="S1")]

	def test_everything_ticked_is_the_slip_net_pay(self):
		# 40000 + 30000 - 4400; the offset pair is hidden and the Calculation Only
		# row is already out of earnings.
		self.assertEqual(
			component_net_amounts(
				self.slips, self.details,
				["Basic Salary", "Festival Allowance"], ["Employee's Contribution SSF"],
			),
			{"S1": 65600.0},
		)

	def test_unticked_earning_is_not_paid(self):
		self.assertEqual(
			component_net_amounts(
				self.slips, self.details, ["Basic Salary"], ["Employee's Contribution SSF"]
			),
			{"S1": 35600.0},
		)

	def test_calculation_only_row_is_never_taken_off(self):
		self.assertEqual(
			component_net_amounts(
				self.slips, self.details, ["Basic Salary"], ["Leave and Late Deduction"]
			),
			{"S1": 40000.0},
		)

	def test_offset_pair_hidden_on_both_sides(self):
		self.assertEqual(
			offset_deduction_roots(self.details), {"Employer's Contribution SSF Deduction"}
		)


class TestBankCodes(TestCase):
	def test_name_variants_find_the_same_bank(self):
		for name in ("Nabil Bank", "Nabil Bank Limited", "NABIL BANK LTD.", " nabil  bank "):
			self.assertEqual(bank_code_for(name, CODES), "0401", name)

	def test_unknown_bank_has_no_code(self):
		self.assertEqual(bank_code_for("Some Other Bank", CODES), "")
		self.assertEqual(bank_code_for(None, CODES), "")


class TestBankRows(TestCase):
	def test_row_layout_matches_the_template_columns(self):
		rows, problems = bank_rows([slip("S1")], {"S1": 1234.567}, CODES, "Salary Ashwin 2083")
		self.assertEqual(problems, [])
		self.assertEqual(rows, [[1234.57, "0401", "0010017501", "A Person", "Salary Ashwin 2083", "EMP-1"]])

	def test_remarks_are_cut_to_the_bank_limit(self):
		rows, _ = bank_rows([slip("S1")], {"S1": 1}, CODES, "Festival Allowance + Bonus Ashwin 2083")
		self.assertEqual(len(rows[0][4]), 30)

	def test_nothing_to_pay_and_cash_paid_are_left_out(self):
		rows, problems = bank_rows(
			[slip("S1"), slip("S2", account="")], {"S1": 0, "S2": 500}, CODES, "Salary"
		)
		self.assertEqual((rows, problems), ([], []))

	def test_missing_bank_code_still_writes_the_row_blank(self):
		rows, problems = bank_rows([slip("S1", bank="Unknown Bank")], {"S1": 10}, CODES, "Salary")
		self.assertEqual(problems, [])
		self.assertEqual(rows[0][:3], [10, "", "0010017501"])

	def test_no_codes_on_the_site_at_all(self):
		rows, problems = bank_rows([slip("S1"), slip("S2")], {"S1": 10, "S2": 20}, {}, "Salary")
		self.assertEqual((len(rows), problems), (2, []))
		self.assertEqual({row[1] for row in rows}, {""})

	def test_negative_amount_is_still_reported(self):
		rows, problems = bank_rows([slip("S1"), slip("S2")], {"S1": 10, "S2": -5}, CODES, "Salary")
		self.assertEqual(len(rows), 1)
		self.assertEqual([p[0].name for p in problems], ["S2"])


class TestAccountFallback(TestCase):
	def test_slip_account_wins(self):
		slips = [slip("S1", account="111")]
		fill_missing_accounts(slips, {"EMP-1": [_dict(bank="Other Bank", bank_account_no="999")]}, {})
		self.assertEqual(slips[0].bank_account_no, "111")

	def test_bank_account_record_fills_a_blank_slip(self):
		slips = [slip("S1", account="", bank=None)]
		fill_missing_accounts(
			slips,
			{"EMP-1": [_dict(bank="Nabil Bank Limited", bank_account_no="06010017500001", custom_nchl_bank_code="")]},
			{"EMP-1": _dict(bank_name="Ignored", bank_ac_no="555")},
		)
		self.assertEqual((slips[0].bank_name, slips[0].bank_account_no), ("Nabil Bank Limited", "06010017500001"))
		rows, problems = bank_rows(slips, {"S1": 100}, CODES, "Salary")
		self.assertEqual((rows[0][1], rows[0][2], problems), ("0401", "06010017500001", []))

	def test_employee_fields_are_the_last_resort(self):
		slips = [slip("S1", account="", bank=None)]
		fill_missing_accounts(slips, {}, {"EMP-1": _dict(bank_name="Nabil Bank", bank_ac_no="555")})
		self.assertEqual((slips[0].bank_name, slips[0].bank_account_no), ("Nabil Bank", "555"))

	def test_no_account_anywhere_stays_blank(self):
		slips = [slip("S1", account="", bank=None)]
		fill_missing_accounts(slips, {"EMP-1": [_dict(bank="Nabil Bank Limited", bank_account_no="")]}, {})
		self.assertEqual(slips[0].bank_account_no, "")


class TestBankWorkbook(TestCase):
	def test_rows_land_under_the_template_headers_as_text(self):
		content = _bank_xlsx([[1234.5, "0401", "0010017501", "A Person", "Salary", "EMP-1"]])
		ws = load_workbook(BytesIO(content)).active
		self.assertEqual(ws.title, "CIPS BULK TRANSFER")
		self.assertEqual(ws.cell(row=4, column=1).value, "Transaction Amount (*)")
		row = [ws.cell(row=BANK_FIRST_ROW, column=c).value for c in range(1, 7)]
		self.assertEqual(row, [1234.5, "0401", "0010017501", "A Person", "Salary", "EMP-1"])
		self.assertEqual(ws.cell(row=BANK_FIRST_ROW, column=2).number_format, "@")
		# A leading 0 is kept and flagged the way Excel flags a typed '0...
		self.assertTrue(ws.cell(row=BANK_FIRST_ROW, column=2).quotePrefix)
		self.assertTrue(ws.cell(row=BANK_FIRST_ROW, column=3).quotePrefix)
		self.assertFalse(ws.cell(row=BANK_FIRST_ROW, column=4).quotePrefix)
		self.assertFalse(ws.cell(row=BANK_FIRST_ROW, column=3).value.startswith("'"))
		self.assertEqual(ws.cell(row=BANK_FIRST_ROW, column=1).number_format, NP_MONEY_FORMAT)


class TestJournalEntryWorkbook(TestCase):
	def test_preview_lists_debits_first_balances_and_sums_expenses(self):
		from nepal_hrms.api.payroll_exports import _jv_xlsx

		pe = _dict(name="PE-1", company="Co", start_date="2026-09-14", end_date="2026-10-14")
		jv = {
			"posted": None,
			"posting_date": "2026-10-15",
			"remark": "Accrual",
			"accounts": [
				_dict(account="Payroll Payable", credit_in_account_currency=900, party_type="Employee", party="E1"),
				_dict(account="Wages", debit_in_account_currency=1000, cost_center="Main"),
				_dict(account="SSF Payable", credit_in_account_currency=100),
			],
		}
		wb = load_workbook(BytesIO(_jv_xlsx(pe, jv, "Ashwin 2083")))
		ws = wb["Journal Entry"]
		self.assertIn("PREVIEW", ws["B2"].value)
		rows = [r for r in ws.iter_rows(min_row=8, values_only=True)]
		self.assertEqual(rows[0][1], "Wages")
		self.assertEqual(rows[-1][1:], ("Total", None, None, None, None, None, 1000, 900 + 100))
		self.assertEqual(list(wb.worksheets[1].iter_rows(values_only=True))[1:], [("Wages", 1000), ("Total", 1000)])


class TestCitPhatwari(TestCase):
	def test_rows_follow_cit_column_order(self):
		from nepal_hrms.api.payroll_exports import cit_rows

		slips = [slip("S1", employee="EMP-1", employee_name="A Person"), slip("S2", employee="EMP-2")]
		ids = {
			"EMP-1": _dict(custom_cit_no="001KBB0000001"),
			"EMP-2": _dict(custom_cit_no="001KBB0000002"),
		}
		self.assertEqual(
			cit_rows(slips, {"S1": 10000, "S2": 0}, ids),
			[["001KBB0000001", "", "", 10000, "A Person"]],
		)

	def test_workbook_is_cit_template_with_text_numbers(self):
		from unittest.mock import patch

		from nepal_hrms.api import payroll_exports as px

		slips = [slip("S1", employee="EMP-1", employee_name="A Person")]
		details = [_dict(parent="S1", parentfield="deductions", salary_component="CIT", amount=5000)]
		ids = {"EMP-1": _dict(custom_cit_no="001KBB0000001")}
		with patch.object(px, "_employee_ids", return_value=ids):
			ws = load_workbook(BytesIO(px._cit_xlsx(None, slips, details, "Ashwin 2083"))).active
		self.assertEqual([c.value for c in ws[1]], ["AC_NO", "OLD_AC_NO", "EMP_SERIAL_NO", "AMOUNT", "CUSTOMER_NAME"])
		self.assertEqual([c.value for c in ws[2]], ["001KBB0000001", None, None, 5000, "A Person"])
		self.assertEqual(ws.max_row, 2)
		self.assertTrue(ws["A2"].quotePrefix)
		self.assertEqual(ws["A2"].number_format, "@")
