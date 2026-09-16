from datetime import date
from unittest import TestCase
from unittest.mock import patch

import frappe

from nepal_hrms.nepal_hrms.report.monthly_attendance_sheet_bs import (
	download_attendance_sheet_bs as exporter,
)
from nepal_hrms.nepal_hrms.report.monthly_attendance_sheet_bs import (
	monthly_attendance_sheet_bs as report,
)


def leave_details():
	return {
		"Home Leave": frappe._dict(
			is_lwp=False, group=report.LEAVE_GROUP_ANNUAL_HOME, configured=True
		),
		"Sick Leave": frappe._dict(
			is_lwp=False, group=report.LEAVE_GROUP_SICK, configured=True
		),
		"Unpaid Leave": frappe._dict(
			is_lwp=True, group=report.LEAVE_GROUP_OTHER, configured=False
		),
		"Training Leave": frappe._dict(
			is_lwp=False, group=report.LEAVE_GROUP_OTHER, configured=False
		),
	}


class TestCanonicalPaidDays(TestCase):
	def classify(self, **kwargs):
		return report.classify_attendance_day(
			leave_type_details=leave_details(), **kwargs
		)

	def test_full_paid_lwp_and_absent_days(self):
		paid = self.classify(status="On Leave", leave_type="Home Leave")
		lwp = self.classify(status="On Leave", leave_type="Unpaid Leave")
		absent = self.classify(status="Absent")

		self.assertEqual(paid["paid_leave_full"], 1)
		self.assertEqual(paid["annual_home_leave"], 1)
		self.assertEqual(paid["paid_flag"], 1)
		self.assertEqual(lwp["lwp_full"], 1)
		self.assertEqual(lwp["paid_flag"], 0)
		self.assertEqual(absent["absent_full"], 1)

	def test_half_days_split_into_two_components(self):
		paid = self.classify(
			status="Half Day",
			half_day_status="Present",
			leave_type="Sick Leave",
		)
		lwp = self.classify(
			status="Half Day",
			half_day_status="Present",
			leave_type="Unpaid Leave",
		)
		cnr = self.classify(status="Half Day", half_day_status="Present")

		self.assertEqual((paid["present_half"], paid["paid_leave_half"]), (0.5, 0.5))
		self.assertEqual(paid["paid_flag"], 1)
		self.assertEqual((lwp["present_half"], lwp["lwp_half"]), (0.5, 0.5))
		self.assertEqual(lwp["paid_flag"], 0.5)
		self.assertEqual((cnr["present_half"], cnr["absent_half"]), (0.5, 0.5))
		self.assertEqual(cnr["paid_flag"], 0.5)

	def test_holiday_work_stays_out_of_normal_present(self):
		day = self.classify(
			status="Present",
			working_hours=7.5,
			is_holiday=True,
		)

		self.assertEqual(day["holiday"], 1)
		self.assertEqual(day["holiday_extra_present"], 1)
		self.assertEqual(day["holiday_extra_hours"], 7.5)
		self.assertEqual(day["present_full"], 0)
		self.assertEqual(day["paid_flag"], 1)

	def test_sandwich_holiday_becomes_paid_lwp_or_absent(self):
		paid = self.classify(
			is_holiday=True,
			is_sandwich=True,
			sandwich_leave_type="Home Leave",
		)
		lwp = self.classify(
			is_holiday=True,
			is_sandwich=True,
			sandwich_leave_type="Unpaid Leave",
		)
		absent = self.classify(is_holiday=True, is_sandwich=True)

		self.assertEqual((paid["holiday"], paid["paid_leave_full"]), (0, 1))
		self.assertEqual((lwp["holiday"], lwp["lwp_full"]), (0, 1))
		self.assertEqual((absent["holiday"], absent["sandwich_absent"]), (0, 1))

	def test_unmarked_day_causes_three_way_mismatch(self):
		row = {
			"total_days_in_period": 2,
			"present_full": 1,
			"present_half": 0,
			"absent_full": 0,
			"absent_half": 0,
			"lwp_full": 0,
			"lwp_half": 0,
			"paid_leave_full": 0,
			"paid_leave_half": 0,
			"holiday": 0,
			"paid_flag": 1,
			"annual_home_leave": 0,
			"sick_leave": 0,
			"other_paid_leave": 0,
		}
		report._finalize_paid_days(row)

		self.assertEqual(row["paid_days_period_method"], 2)
		self.assertEqual(row["paid_days_component_method"], 1)
		self.assertEqual(row["paid_days_daily_method"], 1)
		self.assertEqual(row["payable_days"], 1)
		self.assertEqual(row["payable_days_mismatch"], 1)

	def test_current_period_excludes_future_and_unmarked_today(self):
		current = date(2026, 1, 15)

		self.assertTrue(
			report._include_in_paid_day_audit(
				date(2026, 1, 14), current, has_attendance=False
			)
		)
		self.assertFalse(
			report._include_in_paid_day_audit(
				current, current, has_attendance=False
			)
		)
		self.assertTrue(
			report._include_in_paid_day_audit(
				current, current, has_attendance=True
			)
		)
		self.assertTrue(
			report._include_in_paid_day_audit(
				current, current, is_holiday=True
			)
		)
		self.assertFalse(
			report._include_in_paid_day_audit(
				date(2026, 1, 16), current, has_attendance=True
			)
		)

	def test_occurrence_audit_marks_third_pool_day(self):
		def attendance():
			return frappe._dict(
				employee="EMP-1",
				custom_cnr_late_in_pool=1,
				custom_cnr_early_in_pool=0,
				custom_cnr_unexcused_late=0,
				custom_cnr_unexcused_early=0,
			)
		audit = report._occurrence_audits(
			{
				("EMP-1", date(2026, 1, day)): attendance()
				for day in range(1, 4)
			},
			frappe._dict(
				enable_late_early_penalty=1,
				initial_grace_occurrences=2,
				subsequent_grace_step_occurrences=2,
				combine_late_and_early_pool=0,
				enable_late_full_day_escalation=0,
				late_full_day_after_occurrences=5,
			),
			["EMP-1"],
		)

		self.assertEqual(audit["EMP-1"]["late_half_days"], 1)
		self.assertEqual(audit["EMP-1"]["late_full_days"], 0)
		half = self.classify(status="Present", pool_penalty="half")
		self.assertEqual(half["absent_half"], 0.5)
		self.assertEqual(half["present_half"], 0.5)
		self.assertEqual(half["paid_flag"], 0.5)
		full = self.classify(status="Present", pool_penalty="full")
		self.assertEqual(full["absent_full"], 1)
		self.assertEqual(full["paid_flag"], 0)
		plain_half = self.classify(status="Half Day", half_day_status="Absent")
		self.assertEqual(plain_half["absent_half"], 1.0)
		self.assertEqual(plain_half["paid_flag"], 0)

	def test_summary_column_order_is_exact(self):
		columns = report.get_columns({}, date(2026, 1, 2), date(2026, 1, 1))
		self.assertEqual(
			[column["fieldname"] for column in columns[2:]],
			[
				"total_days_in_period",
				"payable_days",
				"total_absent",
				"lwp_days",
				"paid_leave_days",
				"annual_home_leave_days",
				"sick_leave_days",
				"other_paid_leave_days",
				"holiday_extra_present",
				"holiday_extra_hours",
				"penalty_half_days",
				"penalty_full_days",
				"attendance_request_apps",
				"late_entry_apps",
				"early_exit_apps",
				"out_of_office_apps",
			],
		)


class TestSandwichAndApplications(TestCase):
	@patch.object(report.frappe, "get_all")
	@patch.object(report.frappe.db, "exists")
	def test_batch_sandwich_resolution(self, exists, get_all):
		exists.return_value = True

		def rows_for(doctype, **_kwargs):
			if doctype == "Attendance":
				return [
					frappe._dict(
						employee="EMP-1",
						attendance_date=date(2026, 1, 1),
						status="Absent",
						half_day_status=None,
						leave_type=None,
					),
					frappe._dict(
						employee="EMP-1",
						attendance_date=date(2026, 1, 3),
						status="Absent",
						half_day_status=None,
						leave_type=None,
					),
				]
			if doctype == "Leave Application":
				return [
					frappe._dict(
						employee="EMP-1",
						from_date=date(2026, 1, 1),
						to_date=date(2026, 1, 3),
						leave_type="Home Leave",
						half_day=0,
						half_day_date=None,
						modified="2026-01-01 00:00:00",
					)
				]
			return []

		get_all.side_effect = rows_for
		result = report._build_sandwich_map(
			[frappe._dict(name="EMP-1", holiday_list="HL-1")],
			date(2026, 1, 1),
			date(2026, 1, 3),
			{"HL-1": {date(2026, 1, 2): True}},
			None,
			frappe._dict(enable_leave_sandwich_rule=1),
		)

		self.assertEqual(result[("EMP-1", date(2026, 1, 2))], "Home Leave")

	@patch(
		"nepal_hrms.custom_code.attendance_holiday.covering_leave_application"
	)
	@patch("nepal_hrms.custom_code.attendance_holiday.is_sandwiched")
	def test_sandwich_setting_and_policy_escape_result(self, is_sandwiched, covering):
		settings = frappe._dict(enable_leave_sandwich_rule=0)
		self.assertEqual(
			report._sandwich_day("EMP-1", date(2026, 1, 1), True, settings),
			(False, None),
		)
		is_sandwiched.return_value = False
		settings.enable_leave_sandwich_rule = 1
		self.assertEqual(
			report._sandwich_day("EMP-1", date(2026, 1, 1), True, settings),
			(False, None),
		)
		is_sandwiched.return_value = True
		covering.return_value = frappe._dict(leave_type="Home Leave")
		self.assertEqual(
			report._sandwich_day("EMP-1", date(2026, 1, 1), True, settings),
			(True, "Home Leave"),
		)

	@patch.object(report.frappe, "get_all")
	@patch.object(report.frappe, "get_meta")
	@patch.object(report.frappe.db, "exists")
	def test_application_statuses_and_attendance_request_overlap(
		self, exists, get_meta, get_all
	):
		exists.return_value = True
		get_meta.return_value = frappe._dict(
			has_field=lambda _fieldname: False
		)
		calls = []

		def rows_for(doctype, **kwargs):
			calls.append((doctype, kwargs["filters"]))
			if doctype == "Attendance Request":
				return [
					frappe._dict(
						name="AR-1",
						employee="EMP-1",
						docstatus=1,
						status="Approved",
						from_date=date(2025, 12, 31),
						to_date=date(2026, 1, 2),
					)
				]
			return [
				frappe._dict(
					name=f"{doctype}-P",
					employee="EMP-1",
					docstatus=0,
					status="Open",
					attendance_date=date(2026, 1, 1),
				),
				frappe._dict(
					name=f"{doctype}-R",
					employee="EMP-1",
					docstatus=0,
					status="Rejected",
					attendance_date=date(2026, 1, 1),
				),
			]

		get_all.side_effect = rows_for
		audits, _maps, _day_maps = report._load_application_audits(
			["EMP-1"], date(2026, 1, 1), date(2026, 1, 31)
		)

		self.assertEqual(audits["EMP-1"]["attendance_request"]["approved"], 1)
		self.assertEqual(audits["EMP-1"]["late_entry"]["pending"], 1)
		self.assertEqual(audits["EMP-1"]["late_entry"]["rejected"], 1)
		ar_filters = next(filters for doctype, filters in calls if doctype == "Attendance Request")
		self.assertEqual(ar_filters["from_date"], ["<=", date(2026, 1, 31)])
		self.assertEqual(ar_filters["to_date"], [">=", date(2026, 1, 1)])


class TestAttendanceSheetExport(TestCase):
	def sample_row(self):
		row = {
			"employee": "EMP-1",
			"employee_name": "Employee One",
			"summary": "8:00 / 8:00",
			"application_audits": {},
			"occurrence_audit": {},
			"leave_breakdown": {"Home Leave": 1},
		}
		for fieldname, _label in exporter.SUMMARY_EXPORT_FIELDS:
			row[fieldname] = 1
		row.update(
			{
				"absent_full": 0,
				"absent_half": 0,
				"lwp_full": 0,
				"lwp_half": 0,
				"paid_leave_full": 1,
				"paid_leave_half": 0,
				"paid_days_period_method": 2,
				"paid_days_component_method": 2,
				"paid_days_daily_method": 2,
			}
		)
		return row

	@patch.object(exporter, "_resolve_period")
	@patch.object(exporter, "_day_fieldnames")
	@patch.object(report, "execute")
	def test_daily_details_control_export_columns(
		self, execute, day_fieldnames, resolve_period
	):
		execute.return_value = ([], [self.sample_row()], "")
		resolve_period.return_value = (date(2026, 1, 1), date(2026, 1, 1))
		day_fieldnames.return_value = [
			{
				"fieldname": "day_2026_01_01",
				"label_prefix": "01 Thu (01 Jan 2026)",
			}
		]

		headers, _rows = exporter.build_export_table(
			{"show_details": 0}, export_mode="full"
		)
		self.assertEqual(len(headers), 3 + len(exporter.SUMMARY_EXPORT_FIELDS))
		self.assertNotIn("Absent Full", headers)

		headers, _rows = exporter.build_export_table(
			{"show_details": 1}, export_mode="full"
		)
		self.assertIn("Absent Full", headers)
		self.assertIn("Paid Leave: Home Leave", headers)
		self.assertIn("Paid Days: Canonical Daily Flags", headers)
		self.assertTrue(any("Status" in str(header) for header in headers))
		self.assertTrue(any("Penalty" in str(header) for header in headers))
		self.assertTrue(any("Late Grace" in str(header) for header in headers))

	def test_xlsx_total_columns_are_bold(self):
		from openpyxl import load_workbook
		from io import BytesIO

		headers = ["Employee", "Total Period Days", "Sick Leaves"]
		content = exporter._styled_xlsx_bytes(headers, [["EMP-1", 31, 1]])
		sheet = load_workbook(BytesIO(content)).active

		self.assertTrue(sheet.cell(2, 2).font.bold)
		self.assertFalse(sheet.cell(2, 3).font.bold)
