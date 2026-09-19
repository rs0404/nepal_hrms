# Copyright (c) 2026, Yarsa Labs Pvt. Ltd. and contributors
# For license information, please see license.txt

"""Spreadsheet export for Monthly Attendance Sheet BS.

Standard report export dumps day cells as JSON. These downloads expand each day into
plain columns:
  - daily: Status, Check In, Check Out, Worked Hrs
  - full:  Status, Check In, Check Out, Worked Hrs, Logs, Late/early mins, leave, apps
  - grid:  Classic attendance grid (NAME × BS dates, P / LWP / HDxx codes)
"""

from __future__ import annotations

import json
import re
from copy import copy
from io import BytesIO

import frappe
from frappe import _
from frappe.utils import add_days, cint, flt, getdate

import nepal_compliance.nepali_date_utils.nepali_date as nepali_date

from nepal_hrms.custom_code.leave_type_acronym import leave_type_acronym

BS_MONTH_NAMES = (
	"BAISAKH",
	"JESTHA",
	"ASHAD",
	"SHRAWAN",
	"BHADRA",
	"ASHWIN",
	"KARTIK",
	"MANGSIR",
	"POUSH",
	"MAGH",
	"FALGUN",
	"CHAITRA",
)


SUMMARY_EXPORT_FIELDS = (
	("total_days_in_period", "Total Period Days"),
	("payable_days", "Total Paid Days"),
	("total_absent", "Absent Days"),
	("penalty_days", "Penalty Days"),
	("lwp_days", "Leave Without Pay"),
	("paid_leave_days", "Paid Leaves"),
	("annual_home_leave_days", "Annual/Home Leaves"),
	("sick_leave_days", "Sick Leaves"),
	("other_paid_leave_days", "Other Paid Leaves"),
	("holiday_extra_present", "Present on WO/Hol"),
	("holiday_extra_hours", "Extra Hrs (WO/Hol)"),
	("penalty_half_days", "½P Penalty"),
	("penalty_full_days", "1P Penalty"),
	("attendance_request_apps", "Attendance Request"),
	("late_entry_apps", "Late Entry"),
	("early_exit_apps", "Early Exit"),
	("out_of_office_apps", "Out of Office"),
)

BOLD_TOTAL_EXPORT_FIELDS = {
	"total_days_in_period",
	"payable_days",
	"total_absent",
	"penalty_days",
	"lwp_days",
	"paid_leave_days",
	"holiday_extra_present",
	"holiday_extra_hours",
}

APPLICATION_EXPORT_KINDS = (
	("attendance_request", "Attendance Request"),
	("late_entry", "Late Entry"),
	("early_exit", "Early Exit"),
	("out_of_office", "Out of Office"),
)

GRID_TRAILING_FIELDS = (("summary", "Summary"),) + SUMMARY_EXPORT_FIELDS

GRID_APPLICATION_SUMMARY_FIELDS = {
	"attendance_request_apps",
	"late_entry_apps",
	"early_exit_apps",
	"out_of_office_apps",
}


def _strip_html(label) -> str:
	if label is None:
		return ""
	text = re.sub(r"<[^>]+>", " ", str(label))
	return re.sub(r"\s+", " ", text).strip()


def _parse_day_cell(raw) -> dict:
	if isinstance(raw, dict):
		return raw
	if raw in (None, ""):
		return {}
	if isinstance(raw, (bytes, bytearray)):
		raw = raw.decode("utf-8", errors="ignore")
	try:
		parsed = json.loads(raw)
		return parsed if isinstance(parsed, dict) else {"status": str(raw)}
	except Exception:
		return {"status": str(raw)}


def _normalize_sheet_bs_filters(filters):
	"""Align filter types with the report (Check fields often arrive as \"0\"/\"1\" strings)."""
	filters = frappe._dict(filters or {})
	filters["use_custom_period"] = cint(filters.get("use_custom_period"))
	return filters


def _day_fields_from_report_columns(columns):
	"""Use the same day columns as the on-screen report so export keys always match row data."""
	nepali_date._ensure_loaded()
	fields = []
	for col in columns or []:
		fieldname = col.get("fieldname") or ""
		if not fieldname.startswith("day_"):
			continue
		if fieldname.endswith(("_in", "_out", "_wh", "_checks")):
			continue
		parts = fieldname.replace("day_", "", 1).split("_")
		if len(parts) != 3:
			continue
		try:
			y, m, d = (int(part) for part in parts)
			ad = getdate(f"{y:04d}-{m:02d}-{d:02d}")
		except (TypeError, ValueError):
			continue
		bs = nepali_date.ad_to_bs(ad)
		fields.append(
			{
				"ad": ad,
				"fieldname": fieldname,
				"bs_month": bs["month"],
				"bs_day": bs["day"],
			}
		)
	return fields


def _day_fieldnames(start_ad, end_ad):
	curr = getdate(start_ad)
	end = getdate(end_ad)
	fields = []
	while curr <= end:
		bs = nepali_date.ad_to_bs(curr)
		day = bs["day"]
		weekday = curr.strftime("%a")
		fieldname = f"day_{curr.strftime('%Y_%m_%d')}"
		label_prefix = f"{day:02d} {weekday} ({curr.strftime('%d %b %Y')})"
		fields.append(
			{
				"ad": curr,
				"fieldname": fieldname,
				"label_prefix": label_prefix,
				"bs_label": f"{bs['year']}-{bs['month']:02d}-{bs['day']:02d}",
			}
		)
		curr = add_days(curr, 1)
	return fields


def _resolve_period(filters):
	"""The window the report itself resolved, so the export cannot drift from it.

	This used to be a second copy of the date maths and would have had to learn
	the payroll-period setting independently.
	"""
	from nepal_hrms.nepal_hrms.report.monthly_attendance_sheet_bs.monthly_attendance_sheet_bs import (
		resolve_report_period,
	)

	return resolve_report_period(_normalize_sheet_bs_filters(filters))


def _time_display(value) -> str:
	if value in (None, ""):
		return ""
	text = str(value)
	# "HH:MM:SS" or datetime-like → HH:MM
	if " " in text:
		text = text.split(" ")[-1]
	return text[:5] if len(text) >= 5 else text


def _worked_hours(row: dict, fieldname: str, cell: dict):
	wh = row.get(f"{fieldname}_wh")
	if wh not in (None, ""):
		return round(flt(wh), 2)
	wh_str = cell.get("wh") or ""
	if not wh_str:
		return ""
	# "8h" / "7.5h"
	cleaned = str(wh_str).replace("h", "").strip()
	try:
		return round(float(cleaned), 2)
	except Exception:
		return wh_str


def _day_check_times(row: dict, fieldname: str, cell: dict) -> tuple[str, str]:
	"""Resolve check-in/out for export (detail columns or JSON day cell)."""
	check_in = _time_display(row.get(f"{fieldname}_in") or "")
	check_out = _time_display(row.get(f"{fieldname}_out") or "")
	if not check_in:
		check_in = _time_display(cell.get("in_time") or "")
	if not check_out:
		check_out = _time_display(cell.get("out_time") or "")
	if not check_in and cell.get("checks"):
		parts = [
			part.strip()
			for part in str(cell.get("checks")).split(",")
			if part.strip()
		]
		if parts:
			check_in = parts[0]
			check_out = parts[-1] if len(parts) > 1 else ""
	return check_in, check_out


def _grid_month_title(filters) -> str:
	filters = frappe._dict(filters or {})
	if filters.get("filter_based_on") == "BS Date Range":
		from_bs = (filters.get("from_bs_date") or "").strip()
		try:
			_y, month, _d = map(int, from_bs.split("-"))
			return BS_MONTH_NAMES[max(1, min(month, 12)) - 1]
		except Exception:
			return ""
	try:
		month = int(filters.get("bs_month") or 1)
	except (TypeError, ValueError):
		month = 1
	return BS_MONTH_NAMES[max(1, min(month, 12)) - 1]


def _grid_half_day_suffix(leave_type: str | None, acronym: str | None) -> str:
	"""Two-letter suffix for HDxx codes (matches legacy attendance sheet)."""
	lt = (leave_type or "").strip()
	by_name = {
		"Leave Without Pay": "LWP",
		"Sick Leave": "SL",
		"Annual Leave": "AL",
		"Annual/Home Leave": "HL",
		"Casual Leave": "CL",
	}
	if lt in by_name:
		return by_name[lt]
	ac = (acronym or leave_type_acronym(lt or None) or "Lv").upper()
	if ac == "AHL":
		return "HL"
	if ac == "LWP":
		return "LWP"
	if len(ac) >= 2:
		return ac[:2]
	return ac


def grid_status_code(cell: dict | None) -> str:
	"""Map report day JSON to a classic grid cell code (P, LWP, HDAL, …)."""
	if not cell:
		return ""
	penalty = cell.get("pool_penalty")
	if penalty == "full":
		return str(cell.get("penalty_mark") or "1P")
	if penalty == "half":
		return str(cell.get("penalty_mark") or "½P")

	status = (cell.get("status") or "").strip()
	if status in ("NONE", ""):
		return ""
	if status in ("WO", "Hol"):
		if cint(cell.get("worked_on_holiday")):
			return "P"
		return ""
	if status in ("P", "WFH"):
		return "P"
	if status == "Ab":
		return ""

	leave_type = (cell.get("leave_type") or "").strip()
	acronym = (cell.get("leave_acronym") or leave_type_acronym(leave_type or None) or "").upper()

	if status == "Lv":
		if acronym == "LWP" or leave_type == "Leave Without Pay":
			return "LWP"
		return acronym[:3] if acronym else "Lv"

	if status in ("HD/P", "HD/A"):
		if leave_type or (acronym and acronym not in ("LV",)):
			return f"HD{_grid_half_day_suffix(leave_type, acronym)}"
		if status == "HD/A":
			return "HDA"
		return "P"

	return status


def _grid_summary_display(row: dict, fieldname: str):
	raw = row.get(fieldname)
	if raw in (None, ""):
		return "" if fieldname in GRID_APPLICATION_SUMMARY_FIELDS else 0
	if fieldname in GRID_APPLICATION_SUMMARY_FIELDS or fieldname == "summary":
		return str(raw)
	num = flt(raw)
	rounded = round(num * 10) / 10
	if abs(rounded - round(rounded)) < 1e-9:
		return int(round(rounded))
	return rounded


def build_classic_grid_xlsx_bytes(filters) -> bytes:
	"""Excel attendance grid: MONTH / NAME headers, BS dates, weekend shading."""
	from nepal_hrms.nepal_hrms.report.monthly_attendance_sheet_bs.monthly_attendance_sheet_bs import (
		execute,
	)
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter

	filters = _normalize_sheet_bs_filters(filters)
	columns, data, _message = execute(filters)
	if not data:
		frappe.throw(_("No data found for the selected filters."))

	day_fields = _day_fields_from_report_columns(columns)
	if not day_fields:
		start_ad, end_ad = _resolve_period(filters)
		day_fields = _day_fieldnames(start_ad, end_ad)
	if not day_fields:
		frappe.throw(_("No day columns found for the selected period."))

	month_title = _grid_month_title(filters)
	day_last_col = 1 + len(day_fields)
	summary_start_col = day_last_col + 1
	trailing_fields = list(GRID_TRAILING_FIELDS)
	last_col = summary_start_col + len(trailing_fields) - 1

	wb = Workbook()
	ws = wb.active
	ws.title = "Attendance Sheet BS"

	header_fill = PatternFill("solid", fgColor="B4C7E7")
	# Holiday amber, same as the report pill and the BSHRMS calendar.
	weekend_fill = PatternFill("solid", fgColor="FDE8C4")
	thin = Side(style="thin", color="000000")
	border = Border(left=thin, right=thin, top=thin, bottom=thin)
	header_font = Font(bold=True)
	title_font = Font(bold=True, size=12)
	center = Alignment(horizontal="center", vertical="center")
	left = Alignment(horizontal="left", vertical="center")
	right = Alignment(horizontal="right", vertical="center")

	def style_cell(cell, *, fill=None, bold=False, align=center):
		cell.border = border
		cell.alignment = align
		if fill:
			cell.fill = fill
		if bold:
			cell.font = header_font

	weekend_cols = set()
	for idx, day in enumerate(day_fields, start=2):
		ad = getdate(day["ad"])
		if ad.weekday() >= 5:
			weekend_cols.add(idx)

	# Row 1 — month banner
	ws.cell(1, 1, _("MONTH"))
	style_cell(ws.cell(1, 1), fill=header_fill, bold=True, align=center)
	ws.cell(1, 2, month_title)
	style_cell(ws.cell(1, 2), fill=header_fill, bold=True, align=center)
	title_end_col = max(day_last_col, 3)
	if title_end_col >= 3:
		ws.merge_cells(start_row=1, start_column=3, end_row=1, end_column=title_end_col)
	title_cell = ws.cell(1, 3, _("ATTENDANCE SHEET"))
	style_cell(title_cell, bold=True, align=right)
	title_cell.font = title_font
	if summary_start_col <= last_col:
		ws.merge_cells(
			start_row=1,
			start_column=summary_start_col,
			end_row=1,
			end_column=last_col,
		)
		totals_cell = ws.cell(1, summary_start_col, _("TOTALS"))
		style_cell(totals_cell, fill=header_fill, bold=True, align=center)

	# Row 2 — BS dates
	ws.cell(2, 1, _("NAME"))
	style_cell(ws.cell(2, 1), fill=header_fill, bold=True, align=center)
	for idx, day in enumerate(day_fields, start=2):
		bs_m = day.get("bs_month")
		bs_d = day.get("bs_day")
		if bs_m is None or bs_d is None:
			bs = nepali_date.ad_to_bs(getdate(day["ad"]))
			bs_m, bs_d = bs["month"], bs["day"]
		label = f"{bs_m}/{bs_d}"
		cell = ws.cell(2, idx, label)
		fill = weekend_fill if idx in weekend_cols else header_fill
		style_cell(cell, fill=fill, bold=True, align=center)

	# Row 3 — weekdays
	style_cell(ws.cell(3, 1), fill=header_fill, bold=True, align=center)
	for idx, day in enumerate(day_fields, start=2):
		weekday = getdate(day["ad"]).strftime("%a")
		cell = ws.cell(3, idx, weekday)
		fill = weekend_fill if idx in weekend_cols else header_fill
		style_cell(cell, fill=fill, bold=True, align=center)

	for offset, (_fieldname, label) in enumerate(trailing_fields):
		col = summary_start_col + offset
		header_cell = ws.cell(2, col, _(label))
		style_cell(header_cell, fill=header_fill, bold=True, align=center)
		style_cell(ws.cell(3, col), fill=header_fill, bold=True, align=center)

	# Employee rows
	for row_index, row in enumerate(data, start=4):
		name_cell = ws.cell(row_index, 1, row.get("employee_name") or row.get("employee") or "")
		style_cell(name_cell, align=left)
		for col_index, day in enumerate(day_fields, start=2):
			cell_data = _parse_day_cell(row.get(day["fieldname"]))
			code = grid_status_code(cell_data)
			cell = ws.cell(row_index, col_index, code)
			fill = weekend_fill if col_index in weekend_cols else None
			style_cell(cell, fill=fill, align=center)
		for offset, (fieldname, _label) in enumerate(trailing_fields):
			col = summary_start_col + offset
			value = _grid_summary_display(row, fieldname)
			cell = ws.cell(row_index, col, value)
			style_cell(cell, align=center)

	ws.column_dimensions["A"].width = 24
	for col in range(2, day_last_col + 1):
		ws.column_dimensions[get_column_letter(col)].width = 5.5
	for col in range(summary_start_col, last_col + 1):
		ws.column_dimensions[get_column_letter(col)].width = 14

	output = BytesIO()
	wb.save(output)
	return output.getvalue()


def build_export_table(filters, export_mode="daily"):
	"""Return (headers, rows) ready for Excel/CSV — no JSON day blobs."""
	from nepal_hrms.nepal_hrms.report.monthly_attendance_sheet_bs.monthly_attendance_sheet_bs import (
		execute,
	)

	export_mode = (export_mode or "daily").strip().lower()
	if export_mode not in ("daily", "full", "grid"):
		frappe.throw(_("Invalid export mode. Use daily, full, or grid."))

	filters = _normalize_sheet_bs_filters(filters)
	show_details = bool(cint(filters.get("show_details")))
	include_audit_tail = show_details or export_mode == "full"

	columns, data, _message = execute(filters)
	if not data:
		frappe.throw(_("No data found for the selected filters."))

	day_fields = _day_fields_from_report_columns(columns)
	if not day_fields:
		start_ad, end_ad = _resolve_period(filters)
		day_fields = _day_fieldnames(start_ad, end_ad)

	headers = [_("Employee"), _("Employee Name"), _("Summary")]
	headers.extend(_(label) for _fieldname, label in SUMMARY_EXPORT_FIELDS)

	paid_leave_types = sorted(
		{
			leave_type
			for row in data
			for leave_type in (row.get("leave_breakdown") or {})
		}
	)
	# Spreadsheet download always expands days; not tied to on-screen "Show Daily Details".
	for day in day_fields:
		prefix = day["label_prefix"]
		headers.append(f"{prefix} {_('Status')}")
		headers.append(f"{prefix} {_('Check In')}")
		headers.append(f"{prefix} {_('Check Out')}")
		headers.append(f"{prefix} {_('Worked Hrs')}")
		if show_details:
			headers.append(f"{prefix} {_('Penalty')}")
			headers.append(f"{prefix} {_('Late Grace')}")
			headers.append(f"{prefix} {_('Late Max')}")
			headers.append(f"{prefix} {_('Late Applied')}")
			headers.append(f"{prefix} {_('Late Approved')}")
			headers.append(f"{prefix} {_('Early Grace')}")
			headers.append(f"{prefix} {_('Early Max')}")
			headers.append(f"{prefix} {_('Early Applied')}")
			headers.append(f"{prefix} {_('Early Approved')}")
			headers.append(f"{prefix} {_('OOO Applied')}")
			headers.append(f"{prefix} {_('OOO Approved')}")
		if export_mode == "full":
			headers.extend(
				[
					f"{prefix} {_('Logs')}",
					f"{prefix} {_('Late Mins')}",
					f"{prefix} {_('Early Mins')}",
					f"{prefix} {_('Leave')}",
					f"{prefix} {_('Late App')}",
					f"{prefix} {_('Early App')}",
					f"{prefix} {_('OOO App')}",
				]
			)
	if include_audit_tail:
		headers.extend(
			[
				_("Absent Full"),
				_("Absent Half"),
				_("LWP Full"),
				_("LWP Half"),
				_("Paid Leave Full"),
				_("Paid Leave Half"),
			]
		)
		headers.extend(f"{_('Paid Leave')}: {leave_type}" for leave_type in paid_leave_types)
		headers.extend(
			[
				_("Paid Days: Period - LWP - Absent"),
				_("Paid Days: Present + Paid Leave + Holidays"),
				_("Paid Days: Canonical Daily Flags"),
			]
		)
		for _kind, label in APPLICATION_EXPORT_KINDS:
			headers.extend([f"{_(label)} {_('Pending')}", f"{_(label)} {_('Rejected')}"])
		headers.extend(
			[
				_("Late Pool Occurrences"),
				_("Late Major Occurrences"),
				_("Early Pool Occurrences"),
				_("Early Major Occurrences"),
				_("Initial Grace Occurrences"),
				_("Subsequent Grace Step"),
				_("Free Occurrences Remaining"),
				_("Late Half-Day Equivalents"),
				_("Late Full-Day Equivalents"),
				_("Early Half-Day Equivalents"),
				_("Early Full-Day Equivalents"),
				_("Full-Day Escalation Threshold"),
				_("Combined Late/Early Pool"),
				_("Occurrence Penalty Treatment"),
			]
		)

	out_rows = []
	for row in data:
		line = [
			row.get("employee") or "",
			row.get("employee_name") or "",
			row.get("summary") or "",
		]
		line.extend(row.get(fieldname) for fieldname, _label in SUMMARY_EXPORT_FIELDS)
		for day in day_fields:
			fname = day["fieldname"]
			cell = _parse_day_cell(row.get(fname))
			status = cell.get("status") or ""
			if status == "NONE":
				status = ""
			check_in, check_out = _day_check_times(row, fname, cell)
			worked = _worked_hours(row, fname, cell)
			line.extend([status, check_in, check_out, worked])
			if show_details:
				apps = cell.get("applications") or []
				late_apps = [app for app in apps if app.get("doctype") == "Late Entry Application"]
				early_apps = [app for app in apps if app.get("doctype") == "Early Exit Application"]
				ooo_apps = [app for app in apps if app.get("doctype") == "Out of Office Application"]
				line.extend(
					[
						cell.get("penalty_mark") or "",
						cint(cell.get("late_grace") or 0) or "",
						cint(cell.get("late_max") or 0) or "",
						len(late_apps) or "",
						sum(1 for app in late_apps if app.get("status") == "approved") or "",
						cint(cell.get("early_grace") or 0) or "",
						cint(cell.get("early_max") or 0) or "",
						len(early_apps) or "",
						sum(1 for app in early_apps if app.get("status") == "approved") or "",
						len(ooo_apps) or "",
						sum(1 for app in ooo_apps if app.get("status") == "approved") or "",
					]
				)
			if export_mode == "full":
				line.extend(
					[
						cell.get("checks") or row.get(f"{fname}_checks") or "",
						cint(cell.get("late_mins") or 0) or "",
						cint(cell.get("early_mins") or 0) or "",
						cell.get("leave_type") or "",
						cell.get("late_entry_app") or "",
						cell.get("early_exit_app") or "",
						cell.get("ooo_app") or "",
					]
				)
		if include_audit_tail:
			line.extend(
				[
					row.get("absent_full"),
					row.get("absent_half"),
					row.get("lwp_full"),
					row.get("lwp_half"),
					row.get("paid_leave_full"),
					row.get("paid_leave_half"),
				]
			)
			line.extend(
				(row.get("leave_breakdown") or {}).get(leave_type, 0)
				for leave_type in paid_leave_types
			)
			line.extend(
				[
					row.get("paid_days_period_method"),
					row.get("paid_days_component_method"),
					row.get("paid_days_daily_method"),
				]
			)
			app_audits = row.get("application_audits") or {}
			for kind, _label in APPLICATION_EXPORT_KINDS:
				stats = app_audits.get(kind) or {}
				line.extend([cint(stats.get("pending")), cint(stats.get("rejected"))])
			occurrence = row.get("occurrence_audit") or {}
			free_remaining = occurrence.get("free_remaining")
			if isinstance(free_remaining, dict):
				free_remaining = (
					f"Late {cint(free_remaining.get('late'))}; "
					f"Early {cint(free_remaining.get('early'))}"
				)
			line.extend(
				[
					cint(occurrence.get("late_minor")),
					cint(occurrence.get("late_major")),
					cint(occurrence.get("early_minor")),
					cint(occurrence.get("early_major")),
					cint(occurrence.get("initial_grace_occurrences")),
					cint(occurrence.get("subsequent_grace_step_occurrences")),
					free_remaining if free_remaining is not None else "",
					flt(occurrence.get("late_half_days")),
					flt(occurrence.get("late_full_days")),
					flt(occurrence.get("early_half_days")),
					flt(occurrence.get("early_full_days")),
					cint(occurrence.get("full_day_escalation_threshold")),
					_("Yes") if cint(occurrence.get("combine_pools")) else _("No"),
					_(
						"½P / 1P counted as Absent (gross)"
						if cint(occurrence.get("penalty_enabled"))
						else "Policy disabled; audit projection only"
					),
				]
			)
		out_rows.append(line)

	return headers, out_rows


def _styled_xlsx_bytes(headers, rows):
	from openpyxl import load_workbook
	from frappe.utils.xlsxutils import make_xlsx

	content = make_xlsx([headers] + rows, "Attendance Sheet BS").getvalue()
	workbook = load_workbook(BytesIO(content))
	sheet = workbook.active
	bold_labels = {
		str(_(label))
		for fieldname, label in SUMMARY_EXPORT_FIELDS
		if fieldname in BOLD_TOTAL_EXPORT_FIELDS
	}
	for column_index, header in enumerate(headers, start=1):
		if str(header) not in bold_labels:
			continue
		for row_index in range(2, sheet.max_row + 1):
			cell = sheet.cell(row=row_index, column=column_index)
			font = copy(cell.font)
			font.bold = True
			cell.font = font
	output = BytesIO()
	workbook.save(output)
	return output.getvalue()


@frappe.whitelist()
def download_monthly_attendance_sheet_bs(filters=None, export_mode="daily", file_format="Excel"):
	"""Download Monthly Attendance Sheet BS with expanded daily columns (no JSON)."""
	from frappe.desk.utils import get_csv_bytes, provide_binary_file, pop_csv_params

	if isinstance(filters, str):
		filters = frappe.parse_json(filters) or {}

	frappe.permissions.can_export("Attendance", raise_exception=True)
	filters = _normalize_sheet_bs_filters(filters)

	export_mode = (export_mode or "daily").strip().lower()
	file_format = (file_format or "Excel").strip()

	if export_mode == "grid":
		if file_format.lower() == "csv":
			frappe.throw(_("Classic grid export is available in Excel format only."))
		content = build_classic_grid_xlsx_bytes(filters)
		provide_binary_file("Monthly Attendance Sheet BS (grid)", "xlsx", content)
		return

	headers, rows = build_export_table(filters, export_mode=export_mode)
	xlsx_data = [headers] + rows

	mode_suffix = "daily" if export_mode == "daily" else "full"
	report_name = f"Monthly Attendance Sheet BS ({mode_suffix})"

	if file_format.lower() == "csv":
		csv_params = pop_csv_params(frappe._dict(frappe.local.form_dict))
		content = get_csv_bytes(
			[[frappe.as_unicode(v) if v is not None else "" for v in r] for r in xlsx_data],
			csv_params,
		)
		provide_binary_file(report_name, "csv", content)
	else:
		content = _styled_xlsx_bytes(headers, rows)
		provide_binary_file(report_name, "xlsx", content)
